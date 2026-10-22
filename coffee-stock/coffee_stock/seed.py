"""Начальные данные из файла: справочники, накладные, продажи и инвентаризация.

Сами данные в программу не встроены — они лежат у владельца в файле
«Начальные данные.json». Положите его рядом с CoffeeStock.exe: при запуске
программа перенесёт данные в пустую базу (где ещё нет ни накладных, ни продаж)
и переименует файл в «Начальные данные (загружены).json». Или выберите файл
в «Настройках» → «Загрузить начальные данные…».

Формат:
{
  "ingredients": [{"name": "Кофе в зёрнах", "unit": "г", "min_stock": 2000}, ...],
  "drinks": [{"name": "Капучино", "price": 3.5, "recipe": {"Кофе в зёрнах": 12, ...}}, ...],
  "invoices": [{"number": "...", "date": "2026-06-17", "supplier": "...", "file": "...",
                "lines": [{"name": "...", "qty": 10, "unit": "шт", "price": 40.00, "vat": 80.00,
                           "total": 480.00, "ingredient": "Кофе в зёрнах", "factor": 1000}, ...]}],
  "sales": {"2026-08-01": {"Капучино": 6, ...}, ...},
  "inventory": [{"date": "2026-10-21", "ingredient": "...", "actual": 0, "note": "..."}]
}
Строка накладной с "ingredient": null не идёт в остатки (услуги, моющие средства).
"""

from __future__ import annotations

import json
import sqlite3
from datetime import date, timedelta
from pathlib import Path

from . import core, db, invoice

FILE_NAME = "Начальные данные.json"
LOADED_NAME = "Начальные данные (загружены).json"
SETTING = "seed_applied"


class SeedError(Exception):
    pass


def load(path: Path) -> dict:
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except (OSError, ValueError) as e:
        raise SeedError(f"не удалось прочитать файл: {e}") from e
    if not isinstance(data, dict) or not data.get("ingredients") or not data.get("drinks"):
        raise SeedError("в файле нет списков ингредиентов и напитков")
    names = {i["name"] for i in data["ingredients"]}
    for d in data["drinks"]:
        if unknown := set(d.get("recipe", {})) - names:
            raise SeedError(f"в рецепте «{d['name']}» неизвестные ингредиенты: {', '.join(sorted(unknown))}")
    drinks = {d["name"] for d in data["drinks"]}
    for day, counts in data.get("sales", {}).items():
        if unknown := set(counts) - drinks:
            raise SeedError(f"продажи {day}: неизвестные напитки: {', '.join(sorted(unknown))}")
    return data


def needed(conn: sqlite3.Connection) -> bool:
    """В базе ещё нет ни продаж, ни нормальных накладных — только пример первого запуска.

    Накладные «Не распознана» (например, от старой версии) не мешают: при загрузке они удаляются.
    """
    busy = conn.execute("SELECT (SELECT COUNT(*) FROM invoices WHERE status != 'error') "
                        "+ (SELECT COUNT(*) FROM sales)").fetchone()[0]
    return busy == 0


def apply(conn: sqlite3.Connection, data: dict, sales_file: Path | None = None) -> None:
    with conn:
        for table in ("recipe", "movements", "invoice_lines", "invoices", "aliases", "sales", "drinks",
                      "ingredients"):
            conn.execute(f"DELETE FROM {table}")
        conn.executemany("INSERT INTO ingredients(name, unit, min_stock) VALUES (?, ?, ?)",
                         [(i["name"], i["unit"], i.get("min_stock", 0)) for i in data["ingredients"]])
        ing = {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM ingredients")}
        for d in data["drinks"]:
            drink_id = conn.execute("INSERT INTO drinks(name, price) VALUES (?, ?)",
                                    (d["name"], d.get("price", 0))).lastrowid
            for ing_name, qty in d.get("recipe", {}).items():
                if qty:
                    conn.execute("INSERT INTO recipe VALUES (?, ?, ?)", (drink_id, ing[ing_name], qty))

    for inv in data.get("invoices", []):
        lines = [invoice.EditedLine(l["name"], l["qty"], l.get("unit", ""), l.get("price"), l.get("vat"),
                                    l.get("total"), ing.get(l.get("ingredient")) if l.get("ingredient") else None,
                                    l.get("factor"), not l.get("ingredient"))
                 for l in inv["lines"]]
        invoice_id = invoice.save_invoice(conn, None, inv.get("number", ""), inv.get("date", ""),
                                          inv.get("supplier", ""), lines, post=True)
        if inv.get("file"):
            with conn:
                conn.execute("UPDATE invoices SET file_name = ? WHERE id = ?", (inv["file"], invoice_id))

    with conn:
        drinks = {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM drinks")}
        for day, counts in sorted(data.get("sales", {}).items()):
            core.set_sales_for_date(conn, day, {drinks[n]: c for n, c in counts.items() if c})
        for item in data.get("inventory", []):
            core.adjust_to_actual(conn, item["date"], ing[item["ingredient"]], item["actual"],
                                  item.get("note") or "Инвентаризация")
    db.set_setting(conn, SETTING, "1")
    if sales_file is not None:
        write_sales_file(sales_file, data)
        db.set_setting(conn, "sales_mtime", str(sales_file.stat().st_mtime))


def write_sales_file(path: Path, data: dict) -> None:
    """Файл продаж с уже внесёнными днями и пустыми рабочими датами до конца следующего месяца."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    if path.exists():
        path.replace(path.with_name(path.stem + " (до переноса данных).xlsx"))
    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи"
    names = [d["name"] for d in data["drinks"]]
    ws.append(["Дата"] + names)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="F3E5D0")
    sales = data.get("sales", {})
    for day in sorted(sales):
        ws.append([date.fromisoformat(day)] + [sales[day].get(n) or None for n in names])
    last = date.fromisoformat(max(sales)) if sales else date.today() - timedelta(days=1)
    end = (last.replace(day=1) + timedelta(days=62)).replace(day=1) - timedelta(days=1)
    day = last + timedelta(days=1)
    while day <= end:
        if day.weekday() < 5:
            ws.append([day])
        day += timedelta(days=1)
    for row in ws.iter_rows(min_row=2, max_col=1):
        row[0].number_format = "DD.MM.YYYY"
    ws.column_dimensions["A"].width = 12
    for i in range(len(names)):
        ws.column_dimensions[chr(66 + i)].width = 17
    ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def apply_from_folder(conn: sqlite3.Connection, root: Path, sales_file: Path) -> str:
    """Если рядом с программой лежит файл начальных данных и база пустая — переносит его.

    Возвращает текст для строки состояния (пустой, если делать было нечего).
    """
    path = find_file(root)
    if path is None or not needed(conn):
        return ""
    data = load(path)
    apply(conn, data, sales_file)
    target = root / LOADED_NAME
    if target.exists():
        target.unlink()
    path.rename(target)
    return f"Начальные данные загружены из файла «{path.name}»"


def find_file(root: Path) -> Path | None:
    """Файл начальных данных рядом с программой. Браузер при скачивании мог переименовать его
    в «Начальные данные (1).json» — такие имена тоже подходят."""
    files = sorted((p for p in root.glob("Начальные данные*.json") if p.name != LOADED_NAME),
                   key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0] if files else None
