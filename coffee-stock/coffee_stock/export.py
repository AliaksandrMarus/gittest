"""Выгрузка отчётов в Excel."""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


def save_reports(path: Path, sheets: list[tuple[str, list[str], list[list]]], title: str = "") -> None:
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in sheets:
        ws = wb.create_sheet(name[:31])
        start = 1
        if title:
            ws.cell(row=1, column=1, value=title).font = Font(bold=True, size=12)
            start = 3
        for col, value in enumerate(header, start=1):
            cell = ws.cell(row=start, column=col, value=value)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="F3E5D0")
            ws.column_dimensions[cell.column_letter].width = max(12, len(str(value)) + 4)
        for r, row in enumerate(rows, start=start + 1):
            for col, value in enumerate(row, start=1):
                if isinstance(value, float):
                    value = round(value, 2)
                ws.cell(row=r, column=col, value=value)
        ws.freeze_panes = ws.cell(row=start + 1, column=1)
    wb.save(path)
