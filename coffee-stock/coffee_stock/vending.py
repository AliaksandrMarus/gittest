"""Опрос кофеавтомата Jetinno JL300 по сети (ADB): продажи → «Продажи.xlsx» и отчёты по кофе.

На автомате ничего не меняется — только чтение его базы продаж. Порт отладки 5555 на автомате
включён навсегда. Утилита adb лежит в папке «platform-tools» рядом с программой (если её нет —
скачивается с сайта Google при первом опросе). Ключ, уже разрешённый на автомате, можно положить
в «adb_key\\adbkey»; без него автомат при первом подключении спросит разрешение на своём экране.

Папки (в рабочей папке программы):
  Кофеавтомат\\kofe_ГГГГ-ММ.json — продажи с автомата за месяц (каждая продажа)
  Отчёты кофе\\Кофе ГГГГ-ММ <месяц>.xlsx — отчёт по напиткам и по дням
"""

from __future__ import annotations

import collections
import io
import json
import os
import re
import shutil
import subprocess
import sys
import time
import urllib.request
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

KOFE_DB = "/data/data/com.jinuo.mhwang.jetinnocoffe/databases/coffee.db"
PLATFORM_TOOLS_URL = "https://dl.google.com/android/repository/platform-tools-latest-windows.zip"
MONTHS = ("январь", "февраль", "март", "апрель", "май", "июнь", "июль", "август", "сентябрь", "октябрь",
          "ноябрь", "декабрь")
WEEKDAYS = ("пн", "вт", "ср", "чт", "пт", "сб", "вс")
NO_WINDOW = 0x08000000 if os.name == "nt" else 0
DATA_DIR = "Кофеавтомат"
REPORTS_DIR = "Отчёты кофе"


@dataclass
class PollResult:
    machine: str = ""  # адрес, по которому ответил автомат (мог смениться — найден по MAC)
    months: dict[str, int] = field(default_factory=dict)  # месяц → продаж
    changes: list[str] = field(default_factory=list)  # что внесено в «Продажи.xlsx»
    notes: list[str] = field(default_factory=list)  # предупреждения
    error: str = ""
    reports: list[Path] = field(default_factory=list)


class VendingError(Exception):
    pass


# ---- adb -------------------------------------------------------------------

def adb_path(root: Path) -> Path:
    return root / "platform-tools" / "adb.exe"


def ensure_adb(root: Path) -> Path:
    """adb.exe рядом с программой; если нет — скачать platform-tools с сайта Google."""
    exe = adb_path(root)
    if exe.exists():
        return exe
    if os.name != "nt":
        found = shutil.which("adb")
        if found:
            return Path(found)
        raise VendingError("нет утилиты adb")
    try:
        with urllib.request.urlopen(PLATFORM_TOOLS_URL, timeout=60) as r:
            data = r.read()
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            z.extractall(root)
    except Exception as e:  # noqa: BLE001
        raise VendingError(f"нет утилиты adb и не удалось скачать её с сайта Google: {e}") from e
    if not exe.exists():
        raise VendingError("в скачанном архиве нет adb.exe")
    return exe


def _adb(root: Path, *args: str, timeout: float = 60) -> tuple[int, str, str]:
    exe = adb_path(root) if adb_path(root).exists() else Path(shutil.which("adb") or "adb")
    if getattr(sys, "frozen", False) and os.name == "nt":
        # Фоновый adb-сервер не должен наследовать папку DLL программы — иначе она занята при обновлении
        import ctypes

        ctypes.windll.kernel32.SetDllDirectoryW(None)
    env = {k: v for k, v in os.environ.items() if not k.startswith("_PYI") and k != "_MEIPASS2"}
    key = root / "adb_key" / "adbkey"
    if key.exists():
        env["ADB_VENDOR_KEYS"] = str(key)
    try:
        r = subprocess.run([str(exe), *args], capture_output=True, timeout=timeout, creationflags=NO_WINDOW,
                           env=env, cwd=str(exe.parent) if exe.parent.exists() else None)
    except subprocess.TimeoutExpired:
        return 1, "", "нет ответа (тайм-аут)"
    return r.returncode, r.stdout.decode("utf-8", "replace"), r.stderr.decode("utf-8", "replace")


def kill_server(root: Path) -> None:
    if adb_path(root).exists():
        _adb(root, "kill-server", timeout=10)


def _find_by_mac(mac: str, near: str) -> str | None:
    """IP автомата по MAC: опросить подсеть, чтобы заполнилась таблица ARP."""
    mac = mac.lower().replace(":", "-")
    net = ".".join(near.split(".")[:3])
    procs = [subprocess.Popen(["ping", "-n", "1", "-w", "300", f"{net}.{i}"], stdout=subprocess.DEVNULL,
                              stderr=subprocess.DEVNULL, creationflags=NO_WINDOW) for i in range(1, 255)]
    for p in procs:
        p.wait()
    arp = subprocess.run(["arp", "-a"], capture_output=True, text=True, creationflags=NO_WINDOW).stdout.lower()
    m = re.search(r"(\d+\.\d+\.\d+\.\d+)\s+" + re.escape(mac), arp)
    return m.group(1) if m else None


def connect(root: Path, machine: str, mac: str = "") -> str:
    """Подключиться к автомату → адрес «ip:5555». «offline» — переподключение, сменился IP — поиск по MAC."""
    addr = machine if ":" in machine else f"{machine}:5555"
    for _ in range(2):
        _rc, out, _err = _adb(root, "connect", addr, timeout=20)
        if "connected to" in out:
            _rc, state, err = _adb(root, "-s", addr, "get-state", timeout=15)
            if state.strip() == "device":
                return addr
            if "unauthorized" in state + err:
                raise VendingError("автомат не разрешил доступ этому компьютеру — на экране автомата "
                                   "подтвердите «Разрешить отладку» (галочка «Всегда») и повторите")
        _adb(root, "disconnect", addr, timeout=15)
        time.sleep(2)
    if mac:
        ip = _find_by_mac(mac, addr.split(":")[0])
        if ip and f"{ip}:5555" != addr:
            new = f"{ip}:5555"
            _rc, out, _err = _adb(root, "connect", new, timeout=20)
            if "connected to" in out:
                return new
    raise VendingError(f"кофеавтомат не отвечает ({addr}) — включён ли он и в той же ли сети")


def check(root: Path, machine: str, mac: str = "") -> str:
    """Проверка связи: модель автомата и адрес."""
    ensure_adb(root)
    addr = connect(root, machine, mac)
    _rc, out, _err = _adb(root, "-s", addr, "shell", "getprop ro.product.model", timeout=20)
    _adb(root, "disconnect", addr, timeout=15)
    return f"{addr} — на связи ({out.strip() or 'модель не указана'})"


def fetch_month(root: Path, addr: str, ym: str) -> list[dict]:
    """Продажи за месяц: успешные, без тестовых и бесплатных. Цена на автомате — в копейках."""
    y, m = map(int, ym.split("-"))
    sql = ("select orderNum, productId, productNameEN, price, payType, isOk, timestamp from productInfo "
           f"where timestamp like '{y}{m:02d}%'")
    rc, out, err = _adb(root, "-s", addr, "shell", f"sqlite3 -separator '|' {KOFE_DB} \"{sql}\"", timeout=60)
    if rc != 0 or "rror" in (err or out[:200]):
        raise VendingError(f"не удалось прочитать базу автомата: {(err or out)[:200]}")
    sales = []
    for line in out.splitlines():
        p = line.strip().split("|")
        if len(p) < 7:
            continue
        order, pid, name, price, pay, ok, ts = p[:7]
        if ok.strip() != "1" or pay.strip().lower() in ("test", "free"):
            continue
        try:
            t = datetime.strptime(ts.strip()[:14], "%Y%m%d%H%M%S")
            cost = round(float(price or 0) / 100, 2)
        except ValueError:
            continue
        sales.append({"order": order.strip().lstrip("`"), "date": t.date().isoformat(),
                      "time": t.strftime("%H:%M:%S"), "product_id": pid.strip(), "product": name.strip(),
                      "price": cost, "pay": pay.strip()})
    return sales


# ---- хранение --------------------------------------------------------------

def data_dir(root: Path) -> Path:
    d = root / DATA_DIR
    d.mkdir(parents=True, exist_ok=True)
    return d


def reports_dir(root: Path) -> Path:
    d = root / REPORTS_DIR
    d.mkdir(parents=True, exist_ok=True)
    return d


def save_month(root: Path, ym: str, machine: str, sales: list[dict]) -> None:
    data = {"month": ym, "machine": machine, "updated": datetime.now().isoformat(timespec="seconds"),
            "sales": sales}
    (data_dir(root) / f"kofe_{ym}.json").write_text(json.dumps(data, ensure_ascii=False, indent=1),
                                                     encoding="utf-8")


def load_month(root: Path, ym: str) -> dict | None:
    p = root / DATA_DIR / f"kofe_{ym}.json"
    return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None


def months(root: Path) -> list[str]:
    """Месяцы, за которые есть данные автомата, — новые первыми."""
    folder = root / DATA_DIR
    if not folder.exists():
        return []
    return sorted((p.stem[5:] for p in folder.glob("kofe_*.json")), reverse=True)


def all_sales(root: Path) -> list[dict]:
    out: list[dict] = []
    for ym in months(root):
        out += (load_month(root, ym) or {}).get("sales", [])
    return out


# ---- «Продажи.xlsx» ----------------------------------------------------------

def sync_sales_file(root: Path, path: Path) -> tuple[list[str], list[str]]:
    """Вносит продажи автомата в «Продажи.xlsx» (A — дата, далее столбцы напитков).

    Пустые и недостающие дни заполняются, последний заполненный день обновляется. Если более ранний
    заполненный день расходится с автоматом — только предупреждение (его могли поправить вручную).
    Самая ранняя заполненная строка (например, сводная за прошлый период) не трогается.
    Возвращает (что внесено, предупреждения).
    """
    from openpyxl import load_workbook

    from . import sales_excel

    if not path.exists():
        return [], [f"нет файла {path.name}"]
    try:
        with open(path, "r+b"):
            pass
    except PermissionError:
        return [], [f"«{path.name}» открыт в Excel — продажи внесу, когда файл закроют"]
    wb = load_workbook(path)
    ws = wb[sales_excel.SHEET] if sales_excel.SHEET in wb.sheetnames else wb.active
    col = {str(c.value).strip().lower(): c.column for c in ws[1][1:] if c.value and str(c.value).strip()}
    rows: dict[date, int] = {}
    filled: list[date] = []
    for r in range(2, ws.max_row + 1):
        v = ws.cell(r, 1).value
        if isinstance(v, (datetime, date)):
            d = v.date() if isinstance(v, datetime) else v
            rows[d] = r
            if any(ws.cell(r, c).value not in (None, "", 0) for c in col.values()):
                filled.append(d)
    last_filled = max(filled) if filled else date.min
    first = min(filled) + timedelta(days=1) if filled else date.min
    today = date.today()
    machine: dict[date, collections.Counter] = collections.defaultdict(collections.Counter)
    for s in all_sales(root):
        machine[date.fromisoformat(s["date"])][s["product"].strip().lower()] += 1
    changes: list[str] = []
    notes: list[str] = []
    for d in sorted(x for x in machine if first <= x <= today):
        want = dict(machine[d])
        r = rows.get(d)
        have = {h: int(ws.cell(r, c).value or 0) for h, c in col.items() if ws.cell(r, c).value} if r else {}
        if have == {h: n for h, n in want.items() if h in col}:
            continue
        unknown = [k for k in want if k not in col]
        if unknown:
            notes.append(f"{d:%d.%m}: напитков {', '.join(unknown)} нет в столбцах «{path.name}» — не внесены")
        if have and d < last_filled:
            notes.append(f"{d:%d.%m}: в файле {sum(have.values())} шт, в автомате {sum(want.values())} шт — "
                         "не изменено (проверьте вручную)")
            continue
        if r is None:
            ws.append([datetime(d.year, d.month, d.day)])
            r = ws.max_row
            ws.cell(r, 1).number_format = "DD.MM.YYYY"
            rows[d] = r
        for h, c in col.items():
            ws.cell(r, c).value = want.get(h) or None
        changes.append(f"{d:%d.%m}: {sum(want.values())} шт" + (" (обновлено)" if have else ""))
    if changes:
        backup = root / "Резерв" / f"Продажи {datetime.now():%Y%m%d_%H%M%S}.xlsx"
        backup.parent.mkdir(parents=True, exist_ok=True)
        shutil.copy2(path, backup)
        wb.save(path)
    wb.close()
    return changes, notes


# ---- сводки и отчёты -----------------------------------------------------------

@dataclass
class MonthSummary:
    month: str
    updated: str = ""
    count: int = 0
    money: float = 0.0
    drinks: list[tuple[str, str, int, float, float]] = field(default_factory=list)  # напиток, цены, шт, руб, %
    days: list[tuple[str, str, int, float]] = field(default_factory=list)  # дата, день недели, шт, руб


def summary(root: Path, ym: str) -> MonthSummary:
    data = load_month(root, ym) or {}
    sales = data.get("sales", [])
    s = MonthSummary(ym, data.get("updated", ""), len(sales), round(sum(x["price"] for x in sales), 2))
    by_drink: dict[str, list] = {}
    by_day: dict[str, list] = {}
    for x in sales:
        v = by_drink.setdefault(x["product"], [0, 0.0, collections.Counter()])
        v[0] += 1
        v[1] += x["price"]
        v[2][x["price"]] += 1
        d = by_day.setdefault(x["date"], [0, 0.0])
        d[0] += 1
        d[1] += x["price"]
    for name, (n, total, prices) in sorted(by_drink.items(), key=lambda kv: -kv[1][0]):
        s.drinks.append((name, ", ".join(f"{p:.2f}".replace(".", ",") for p, _ in prices.most_common()), n,
                         round(total, 2), round(100 * n / s.count, 1) if s.count else 0.0))
    for day in sorted(by_day):
        n, total = by_day[day]
        s.days.append((day, WEEKDAYS[date.fromisoformat(day).weekday()], n, round(total, 2)))
    return s


def month_title(ym: str) -> str:
    y, m = map(int, ym.split("-"))
    return f"{MONTHS[m - 1].capitalize()} {y}"


def report_path(root: Path, ym: str) -> Path:
    y, m = map(int, ym.split("-"))
    return reports_dir(root) / f"Кофе {ym} {MONTHS[m - 1]}.xlsx"


def build_report(root: Path, ym: str) -> Path | None:
    """«Кофе ГГГГ-ММ <месяц>.xlsx»: листы «По напиткам», «По дням» (дни × напитки), «Все продажи»."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    data = load_month(root, ym)
    if not data:
        return None
    sales = sorted(data.get("sales", []), key=lambda x: (x["date"], x["time"]))
    s = summary(root, ym)
    order = [d[0] for d in s.drinks]
    updated = s.updated.replace("T", " ")
    wb = Workbook()
    head_fill, total_fill = PatternFill("solid", fgColor="5D4037"), PatternFill("solid", fgColor="F5EBDD")
    thin = Side(style="thin", color="D8C3AC")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)

    def head(ws, title: str, cols: list[str]) -> None:
        ws.append([title])
        ws["A1"].font = Font(bold=True, size=14, color="5D4037")
        ws.append([f"{month_title(ym)} · кофеавтомат · данные автомата на {updated}"])
        ws["A2"].font = Font(italic=True, color="7A6656", size=9)
        ws.append([])
        ws.append(cols)
        for c in range(1, len(cols) + 1):
            cell = ws.cell(4, c)
            cell.font, cell.fill, cell.border = Font(bold=True, color="FFFFFF"), head_fill, box
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.row_dimensions[4].height = 30
        ws.freeze_panes = "A5"

    def total_row(ws, values: list) -> None:
        ws.append(values)
        for c in range(1, len(values) + 1):
            cell = ws.cell(ws.max_row, c)
            cell.font, cell.fill, cell.border = Font(bold=True), total_fill, box

    ws = wb.active
    ws.title = "По напиткам"
    head(ws, "Продажи кофе по напиткам", ["Напиток", "Цена, руб", "Кол-во, шт", "Сумма, руб", "Доля, %"])
    for row in s.drinks:
        ws.append(list(row))
        for c in range(1, 6):
            ws.cell(ws.max_row, c).border = box
    total_row(ws, ["ИТОГО", "", s.count, s.money, 100.0 if s.count else 0])
    for letter, width in zip("ABCDE", (24, 12, 12, 13, 10)):
        ws.column_dimensions[letter].width = width
    for r in range(5, ws.max_row + 1):
        ws.cell(r, 4).number_format = "0.00"

    ws = wb.create_sheet("По дням")
    head(ws, "Продажи кофе по дням", ["Дата", "День"] + order + ["Всего, шт", "Сумма, руб"])
    per: dict[str, collections.Counter] = collections.defaultdict(collections.Counter)
    for x in sales:
        per[x["date"]][x["product"]] += 1
    for day, weekday, n, total in s.days:
        row = [date.fromisoformat(day), weekday] + [per[day].get(k) or None for k in order] + [n, total]
        ws.append(row)
        r = ws.max_row
        ws.cell(r, 1).number_format = "DD.MM.YYYY"
        ws.cell(r, len(row)).number_format = "0.00"
        for c in range(1, len(row) + 1):
            ws.cell(r, c).border = box
            ws.cell(r, c).alignment = Alignment(horizontal="center")
    total_row(ws, ["ИТОГО", ""] + [d[2] for d in s.drinks] + [s.count, s.money])
    ws.column_dimensions["A"].width, ws.column_dimensions["B"].width = 12, 6
    for i in range(len(order) + 2):
        ws.column_dimensions[get_column_letter(3 + i)].width = 13

    ws = wb.create_sheet("Все продажи")
    head(ws, "Все продажи кофеавтомата", ["Дата", "Время", "Напиток", "Цена, руб", "Оплата", "№ заказа"])
    for x in sales:
        ws.append([date.fromisoformat(x["date"]), x["time"], x["product"], x["price"],
                   "карта" if x.get("pay") == "mdb_cashless" else x.get("pay"), x["order"]])
        ws.cell(ws.max_row, 1).number_format = "DD.MM.YYYY"
        ws.cell(ws.max_row, 4).number_format = "0.00"
    total_row(ws, ["ИТОГО", "", f"{s.count} шт", s.money, "", ""])
    for letter, width in zip("ABCDEF", (12, 10, 22, 11, 10, 28)):
        ws.column_dimensions[letter].width = width

    out = report_path(root, ym)
    try:
        wb.save(out)
    except PermissionError:
        return None  # отчёт открыт в Excel — обновится при следующем опросе
    return out


# ---- полный цикл -----------------------------------------------------------------

def poll(root: Path, machine: str, mac: str, sales_file: Path, today: date | None = None) -> PollResult:
    """Опрос автомата (текущий месяц, в первые дни — и прошлый) → «Продажи.xlsx» → отчёты."""
    today = today or date.today()
    result = PollResult()
    wanted = [today.strftime("%Y-%m")]
    prev = (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    if today.day <= 3 or load_month(root, prev) is None:
        wanted.insert(0, prev)
    try:
        ensure_adb(root)
        addr = connect(root, machine, mac)
        result.machine = addr
        try:
            for ym in wanted:
                sales = fetch_month(root, addr, ym)
                save_month(root, ym, addr, sales)
                result.months[ym] = len(sales)
        finally:
            _adb(root, "disconnect", addr, timeout=15)  # старый adbd (Android 6) зависает от висящих соединений
    except VendingError as e:
        result.error = str(e)
        return result
    result.changes, result.notes = sync_sales_file(root, sales_file)
    for ym in wanted:
        if (path := build_report(root, ym)) is not None:
            result.reports.append(path)
        else:
            result.notes.append(f"отчёт «{report_path(root, ym).name}» открыт в Excel — не обновлён")
    return result
