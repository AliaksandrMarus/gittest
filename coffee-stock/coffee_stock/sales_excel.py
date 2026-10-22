"""Excel-файл продаж: строки — даты, столбцы — напитки."""

from __future__ import annotations

import sqlite3
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill

from . import core

SHEET = "Продажи"


@dataclass
class ImportResult:
    days: list[str] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)


def ensure_template(conn: sqlite3.Connection, path: Path) -> None:
    """Создаёт файл продаж или дописывает в него столбцы для новых напитков."""
    drinks = [r["name"] for r in conn.execute("SELECT name FROM drinks ORDER BY id")]
    if path.exists():
        wb = load_workbook(path)
        ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
        have = {str(c.value).strip() for c in ws[1] if c.value}
        missing = [name for name in drinks if name not in have]
        for col, name in enumerate(missing, start=ws.max_column + 1):
            _header_cell(ws, col, name)
        if missing:  # без изменений не сохраняем — иначе зря меняется дата файла
            wb.save(path)
        return

    wb = Workbook()
    ws = wb.active
    ws.title = SHEET
    _header_cell(ws, 1, "Дата")
    for i, name in enumerate(drinks, start=2):
        _header_cell(ws, i, name)
    # Даты на месяц вперёд, чтобы оставалось только вписывать количество
    start = date.today()
    for i in range(31):
        cell = ws.cell(row=i + 2, column=1, value=start + timedelta(days=i))
        cell.number_format = "DD.MM.YYYY"
    ws.column_dimensions["A"].width = 12
    ws.freeze_panes = "B2"
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _header_cell(ws, col: int, value: str) -> None:
    cell = ws.cell(row=1, column=col, value=value)
    cell.font = Font(bold=True)
    cell.fill = PatternFill("solid", fgColor="F3E5D0")
    ws.column_dimensions[cell.column_letter].width = max(12, len(value) + 2)


def _parse_date(value) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in ("%d.%m.%Y", "%d.%m.%y", "%Y-%m-%d", "%d/%m/%Y"):
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                pass
    return None


def import_file(conn: sqlite3.Connection, path: Path) -> ImportResult:
    """Загружает продажи из Excel. Пустые строки (день ещё не заполнен) пропускаются."""
    result = ImportResult()
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[SHEET] if SHEET in wb.sheetnames else wb.active
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None)
    if not header:
        result.errors.append("Файл пустой")
        return result

    drinks = {r["name"].strip().lower(): r["id"] for r in conn.execute("SELECT id, name FROM drinks")}
    columns: dict[int, int] = {}
    for idx, title in enumerate(header[1:], start=1):
        if title is None or not str(title).strip():
            continue
        drink_id = drinks.get(str(title).strip().lower())
        if drink_id is None:
            result.errors.append(f"Напиток «{title}» не найден в рецептах — столбец пропущен")
        else:
            columns[idx] = drink_id

    parsed: dict[str, dict[int, int]] = {}
    for line_no, row in enumerate(rows, start=2):
        if not row or all(v in (None, "") for v in row[1:]):
            continue
        day = _parse_date(row[0])
        if day is None:
            result.errors.append(f"Строка {line_no}: не понял дату «{row[0]}»")
            continue
        counts: dict[int, int] = {}
        for idx, drink_id in columns.items():
            value = row[idx] if idx < len(row) else None
            if value in (None, ""):
                continue
            try:
                n = int(round(float(str(value).replace(",", "."))))
            except ValueError:
                result.errors.append(f"Строка {line_no}: «{value}» — не число")
                continue
            if n < 0:
                result.errors.append(f"Строка {line_no}: отрицательное количество {n}")
                continue
            counts[drink_id] = counts.get(drink_id, 0) + n
        parsed[day] = counts
    wb.close()

    with conn:
        for day, counts in sorted(parsed.items()):
            core.set_sales_for_date(conn, day, counts)
            result.days.append(day)
    return result


# ---- Загрузка Excel-файла пользователя ---------------------------------------
#
# Поддерживаются три вида листов:
#   «таблица»  — строки — дни, столбцы — напитки (или наоборот): количество за день
#                или показания счётчиков автомата нарастающим итогом;
#   «отчёт»    — «Отчёт по реализации»: Дата | Наименование | кол-во, дата только
#                в первой строке дня;
#   наш шаблон «Продажи.xlsx» — частный случай «таблицы».
# Показания нарастающим итогом превращаются в продажи за день как разница
# с предыдущим показанием (в том числе с последним показанием прошлого листа).

MONTHS = ("январ", "феврал", "март", "апрел", "ма", "июн", "июл", "август", "сентябр", "октябр", "ноябр",
          "декабр")
TOTAL_WORDS = ("итог", "всего", "сумма за")
QTY_WORDS = ("кол", "шт", "показ", "счет", "счёт", "продано", "порц")


@dataclass
class SheetSales:
    title: str
    kind: str  # "таблица" или "отчёт"
    values: dict[str, dict[int, float]]  # день → напиток → число (как в файле)
    cumulative: bool = False  # значения — показания счётчиков нарастающим итогом


@dataclass
class WorkbookSales:
    sheets: list[SheetSales] = field(default_factory=list)
    errors: list[str] = field(default_factory=list)

    @property
    def cumulative(self) -> bool:
        return any(s.cumulative for s in self.sheets)

    def daily(self, cumulative: bool | None = None) -> tuple[dict[str, dict[int, int]], list[str]]:
        """Продажи по дням. cumulative=None — как определено по листам; True/False — для всех листов.

        Возвращает продажи и замечания (сброс счётчика, первое показание).
        """
        notes: list[str] = []
        plain: dict[str, dict[int, float]] = {}
        readings: dict[str, dict[int, float]] = {}
        for sheet in self.sheets:
            target = readings if (sheet.cumulative if cumulative is None else cumulative) else plain
            for day, counts in sheet.values.items():
                target.setdefault(day, {}).update(counts)
        result = {day: {d: int(round(v)) for d, v in counts.items() if round(v) > 0}
                  for day, counts in plain.items()}
        last: dict[int, tuple[str, float]] = {}
        for day in sorted(readings):
            counts: dict[int, int] = {}
            known = False
            for drink, value in readings[day].items():
                prev = last.get(drink)
                last[drink] = (day, value)
                if prev is None:
                    continue
                known = True
                diff = value - prev[1]
                if diff < 0:
                    notes.append(f"{fmt_day(day)}: показание меньше предыдущего ({fmt_num(prev[1])} → "
                                 f"{fmt_num(value)}) — счётчик сброшен? День для этого напитка пропущен")
                    continue
                if round(diff) > 0:
                    counts[drink] = int(round(diff))
            if known:
                result.setdefault(day, {}).update(counts)
        if readings:
            notes.append(f"Первое показание {fmt_day(min(readings))} — точка отсчёта, продажи считаются "
                         "со следующего дня")
        return result, notes


def fmt_day(day: str) -> str:
    return date.fromisoformat(day).strftime("%d.%m.%Y")


def fmt_num(value: float) -> str:
    return str(int(value)) if float(value).is_integer() else str(value)


def _norm(text) -> str:
    return " ".join(str(text).lower().replace("ё", "е").split()) if text is not None else ""


class DrinkMatcher:
    def __init__(self, conn: sqlite3.Connection):
        import difflib

        self._ratio = lambda a, b: difflib.SequenceMatcher(None, a, b).ratio()
        self.drinks = {_norm(r["name"]): r["id"] for r in conn.execute("SELECT id, name FROM drinks")}

    def __call__(self, text) -> int | None:
        name = _norm(text)
        if len(name) < 3 or not any(ch.isalpha() for ch in name):
            return None
        if name in self.drinks:
            return self.drinks[name]
        best, best_ratio = None, 0.0
        for known, drink_id in self.drinks.items():
            ratio = self._ratio(name, known)
            if ratio > best_ratio:
                best, best_ratio = drink_id, ratio
        return best if best_ratio >= 0.85 else None


def _number(value) -> float | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip().replace(" ", "").replace("\xa0", "").replace(",", "."))
        except ValueError:
            return None
    return None


def _grid(ws) -> list[list]:
    """Значения листа; объединённые ячейки заполнены значением левой верхней."""
    grid = [list(row) for row in ws.iter_rows(values_only=True)]
    for merged in ws.merged_cells.ranges:
        value = ws.cell(merged.min_row, merged.min_col).value
        for r in range(merged.min_row - 1, min(merged.max_row, len(grid))):
            for c in range(merged.min_col - 1, merged.max_col):
                if c < len(grid[r]):
                    grid[r][c] = value
    width = max((len(r) for r in grid), default=0)
    return [r + [None] * (width - len(r)) for r in grid]


def _month_of(*texts) -> tuple[int, int] | None:
    """Месяц и год из названия листа или заголовка: «октябрь 26», «октябрь 2026 г.»."""
    for text in texts:
        t = _norm(text)
        year = re.search(r"(20\d\d|\b\d\d\b)", t)
        for i, stem in enumerate(MONTHS, start=1):
            pattern = r"\bма[йяе]\b" if stem == "ма" else r"\b" + stem
            if re.search(pattern, t) and year:
                y = int(year.group(1))
                return i, y + 2000 if y < 100 else y
    return None


def _is_total(row: list) -> bool:
    return any(isinstance(v, str) and any(w in _norm(v) for w in TOTAL_WORDS) for v in row)


def _parse_report(grid: list[list], match: DrinkMatcher, title: str, errors: list[str]) -> SheetSales | None:
    """«Отчёт по реализации»: столбцы «Дата», «Наименование», «кол-во»."""
    for h, row in enumerate(grid[:40]):
        cells = [_norm(v) for v in row]
        date_col = next((i for i, v in enumerate(cells) if v == "дата"), None)
        name_col = next((i for i, v in enumerate(cells) if v.startswith("наименование")), None)
        if date_col is None or name_col is None:
            continue
        qty_col = None
        for r in grid[h:h + 2]:
            qty_col = next((i for i, v in enumerate(r) if i not in (date_col, name_col)
                            and any(_norm(v).startswith(w) for w in ("кол", "количество"))), qty_col)
        if qty_col is None:
            continue
        values: dict[str, dict[int, float]] = {}
        day = None
        unknown: set[str] = set()
        for row in grid[h + 1:]:
            day = _parse_date(row[date_col]) or day
            name, qty = row[name_col], _number(row[qty_col])
            if name is None or qty is None or _is_total(row):
                continue
            drink = match(name)
            if drink is None:
                unknown.add(str(name).strip())
                continue
            if day is None or day > date.today().isoformat():
                continue
            values.setdefault(day, {})
            values[day][drink] = values[day].get(drink, 0) + qty
        for name in sorted(unknown):
            errors.append(f"Лист «{title}»: «{name}» нет среди напитков — пропущено")
        return SheetSales(title, "отчёт", values) if values else None
    return None


def _drink_columns(grid: list[list], h: int, match: DrinkMatcher) -> dict[int, int]:
    """Столбцы с количеством для каждого напитка из строки заголовка h.

    Если название напитка объединено над несколькими столбцами («кол-во | сумма»),
    берётся столбец с подписью вроде «кол-во», иначе первый.
    """
    spans: dict[int, list[int]] = {}
    for c, v in enumerate(grid[h]):
        drink = match(v) if isinstance(v, str) else None
        if drink is not None:
            spans.setdefault(drink, []).append(c)
    columns: dict[int, int] = {}
    sub = grid[h + 1] if h + 1 < len(grid) else [None] * len(grid[h])
    for drink, cols in spans.items():
        pick = next((c for c in cols if any(w in _norm(sub[c]) for w in QTY_WORDS)), cols[0])
        columns[pick] = drink
    return columns


DAY_TITLES = ("дата", "число", "день", "дни")


def _day_column(grid: list[list], start: int, skip: set[int], month: tuple[int, int] | None):
    """Столбец с датами (или с числами месяца 1–31, если месяц известен). Возвращает (столбец, функция)."""
    best = None
    for c in range(len(grid[0]) if grid else 0):
        if c in skip:
            continue
        dates = sum(1 for r in grid[start:] if _parse_date(r[c]))
        titled = start > 0 and _norm(grid[start - 1][c]) in DAY_TITLES
        if (dates >= 2 or (dates and titled)) and (best is None or dates > best[0]):
            best = (dates, c, lambda v: _parse_date(v))
    if best:
        return best[1], best[2]
    if month:
        m, y = month

        def day_of(v):
            n = _number(v)
            if n is None or not n.is_integer() or not 1 <= n <= 31:
                return None
            try:
                return date(y, m, int(n)).isoformat()
            except ValueError:
                return None
        for c in range(len(grid[0]) if grid else 0):
            if c in skip:
                continue
            days = [day_of(r[c]) for r in grid[start:]]
            seq = [d for d in days if d]
            if len(seq) >= 2 and seq == sorted(seq) and len(set(seq)) == len(seq):
                return c, day_of
    return None, None


def _looks_cumulative(values: dict[str, dict[int, float]]) -> bool:
    """Показания нарастающим итогом: у каждого напитка не убывают и велики по сравнению с приростом."""
    days = sorted(values)
    if len(days) < 3:
        return False
    series: dict[int, list[float]] = {}
    for day in days:
        for drink, v in values[day].items():
            series.setdefault(drink, []).append(v)
    growing = 0
    for seq in series.values():
        if len(seq) < 3:
            continue
        steps = [b - a for a, b in zip(seq, seq[1:])]
        if any(s < 0 for s in steps):
            return False
        if seq[0] > 0 and seq[-1] >= seq[0] + max(steps) and seq[0] >= max(steps):
            growing += 1
    return growing > 0 and growing >= len([s for s in series.values() if len(s) >= 3]) / 2


def _parse_table(grid: list[list], match: DrinkMatcher, title: str, errors: list[str]) -> SheetSales | None:
    month = _month_of(title, *[v for r in grid[:12] for v in r if isinstance(v, str)])
    for h, row in enumerate(grid[:40]):
        columns = _drink_columns(grid, h, match)
        if len(columns) < 2:
            continue
        day_col, day_of = _day_column(grid, h + 1, set(columns), month)
        if day_col is None:
            continue
        values: dict[str, dict[int, float]] = {}
        for r in grid[h + 1:]:
            if _is_total(r):
                continue
            day = day_of(r[day_col])
            if day is None or day > date.today().isoformat():
                continue
            counts = {drink: n for c, drink in columns.items() if (n := _number(r[c])) is not None}
            # Будущие дни месяца: пусто или одни нули после последнего заполненного дня
            if counts:
                values[day] = counts
        while values and not any(values[max(values)].values()):
            del values[max(values)]
        if values:
            return SheetSales(title, "таблица", values)
    return None


def _transpose(grid: list[list]) -> list[list]:
    return [list(col) for col in zip(*grid)] if grid else []


def read_workbook(conn: sqlite3.Connection, path: Path) -> WorkbookSales:
    """Читает продажи из всех листов Excel-файла (форматы — см. выше). В базу ничего не пишет."""
    result = WorkbookSales()
    match = DrinkMatcher(conn)
    wb = load_workbook(path, data_only=True)
    try:
        for ws in wb.worksheets:
            if ws.sheet_state != "visible":
                continue
            grid = _grid(ws)
            if not grid:
                continue
            sheet = (_parse_report(grid, match, ws.title, result.errors)
                     or _parse_table(grid, match, ws.title, result.errors)
                     or _parse_table(_transpose(grid), match, ws.title, result.errors))
            if sheet:
                result.sheets.append(sheet)
    finally:
        wb.close()
    # Показания нарастающим итогом определяются по всем листам-таблицам вместе:
    # в начале месяца на листе может быть всего пара дней
    tables = [s for s in result.sheets if s.kind == "таблица"]
    combined: dict[str, dict[int, float]] = {}
    for sheet in tables:
        for day, counts in sheet.values.items():
            combined.setdefault(day, {}).update(counts)
    if _looks_cumulative(combined):
        for sheet in tables:
            sheet.cumulative = True
    if not result.sheets:
        result.errors.append("Не нашёл в файле продаж: нужна таблица с названиями напитков и датами "
                             "(или «Дата | Наименование | кол-во»)")
    return result


def apply_daily(conn: sqlite3.Connection, daily: dict[str, dict[int, int]]) -> list[str]:
    with conn:
        for day, counts in sorted(daily.items()):
            core.set_sales_for_date(conn, day, counts)
    return sorted(daily)


def import_workbook(conn: sqlite3.Connection, path: Path, cumulative: bool | None = None) -> ImportResult:
    """Читает файл и сразу записывает продажи (для автоматической перезагрузки изменённого файла)."""
    data = read_workbook(conn, path)
    daily, notes = data.daily(cumulative)
    result = ImportResult(errors=data.errors + [n for n in notes if not n.startswith("Первое показание")])
    result.days = apply_daily(conn, daily)
    return result
