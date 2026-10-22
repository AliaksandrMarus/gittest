"""Загрузка Excel-файлов продаж пользователя: отчёт по реализации, таблица по дням, счётчики нарастающим итогом."""

from datetime import date, datetime, timedelta

import pytest
from openpyxl import Workbook

from coffee_stock import db, sales_excel


@pytest.fixture
def conn(tmp_path):
    return db.connect(tmp_path / "t.db")


def ids(conn):
    return {r["name"]: r["id"] for r in conn.execute("SELECT id, name FROM drinks")}


def by_name(conn, daily):
    names = {v: k for k, v in ids(conn).items()}
    return {day: {names[d]: n for d, n in counts.items()} for day, counts in daily.items()}


def sales_in_db(conn, day):
    return {r[0]: r[1] for r in conn.execute(
        "SELECT d.name, s.count FROM sales s JOIN drinks d ON d.id = s.drink_id WHERE s.date = ?", (day,))}


def test_realization_report_inherits_date_and_ignores_case(conn, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "2025(9)"
    ws.append(["Отчет по реализации товаров"])
    ws.append([])
    ws.append(["№ п/п", "Дата", "Наименование", None, "Единица измерения документа", None, "Отпускная цена",
               "Сумма всего"])
    ws.append([None, None, None, None, "наим-ние", "кол-во"])
    ws.append([1, datetime(2025, 9, 1), "Капучино", None, "шт", 16, 3.5, 56])
    ws.append([2, None, "  горячий   ШОКОЛАД ", None, "шт", 1, 3, 3])
    ws.append([3, None, "Пирожок", None, "шт", 2, 1, 2])
    ws.append([4, datetime(2025, 9, 2), "Эспрессо ", None, "шт", 3, 2.5, 7.5])
    ws.append([None, None, "ИТОГО", None, None, 22, None, 68.5])
    path = tmp_path / "отчёт.xlsx"
    wb.save(path)

    data = sales_excel.read_workbook(conn, path)
    assert [(s.kind, s.cumulative) for s in data.sheets] == [("отчёт", False)]
    assert any("Пирожок" in e for e in data.errors)
    daily, _ = data.daily()
    assert by_name(conn, daily) == {"2025-09-01": {"Капучино": 16, "Горячий шоколад": 1},
                                    "2025-09-02": {"Эспрессо": 3}}


def month_sheet(wb, title, first_day, rows, merged=True):
    """Лист «как у пользователя»: числа месяца в строках, у каждого напитка столбцы «кол-во | сумма»."""
    ws = wb.create_sheet(title)
    ws.append(["Продажи кофейного автомата", None, None, None, None, None, None])
    header = ["Число"]
    for name in ("Капучино", "Американо", "Латте"):
        header += [name, name if not merged else None]
    ws.append(header)
    if merged:
        for col in (2, 4, 6):
            ws.merge_cells(start_row=2, start_column=col, end_row=2, end_column=col + 1)
    ws.append([None] + ["кол-во", "сумма"] * 3)
    for i, counts in enumerate(rows):
        row = [first_day + i]
        for n in counts:
            row += [n, None if n is None else n * 3]
        ws.append(row)
    ws.append(["Итого"] + [999] * 6)
    return ws


def test_month_table_with_day_numbers_and_merged_headers(conn, tmp_path):
    wb = Workbook()
    wb.remove(wb.active)
    month_sheet(wb, "сентябрь 25", 1, [(5, 2, 0), (7, None, 1), (None, None, None)])
    path = tmp_path / "ОБЩИЙ.xlsx"
    wb.save(path)

    data = sales_excel.read_workbook(conn, path)
    assert [(s.kind, s.cumulative) for s in data.sheets] == [("таблица", False)]
    daily, _ = data.daily()
    assert by_name(conn, daily) == {"2025-09-01": {"Капучино": 5, "Американо": 2},
                                    "2025-09-02": {"Капучино": 7, "Латте": 1}}


def test_cumulative_counters_across_month_sheets(conn, tmp_path):
    """Показания счётчиков растут; продажи — разница, в том числе через границу месяца."""
    wb = Workbook()
    wb.remove(wb.active)
    month_sheet(wb, "сентябрь 25", 28, [(100, 50, 10), (105, 52, 10), (111, 52, 12)])
    month_sheet(wb, "октябрь 25", 1, [(120, 55, 12), (124, 55, 13), (None, None, None)])
    path = tmp_path / "ОБЩИЙ.xlsx"
    wb.save(path)

    data = sales_excel.read_workbook(conn, path)
    assert all(s.cumulative for s in data.sheets)
    daily, notes = data.daily()
    assert by_name(conn, daily) == {
        "2025-09-29": {"Капучино": 5, "Американо": 2},
        "2025-09-30": {"Капучино": 6, "Латте": 2},
        "2025-10-01": {"Капучино": 9, "Американо": 3},
        "2025-10-02": {"Капучино": 4, "Латте": 1},
    }
    assert "2025-09-28" not in daily  # первое показание — только точка отсчёта
    assert any("точка отсчёта" in n for n in notes)

    # Пользователь может сказать, что это не показания, а количество за день
    plain, _ = data.daily(cumulative=False)
    assert by_name(conn, plain)["2025-09-28"] == {"Капучино": 100, "Американо": 50, "Латте": 10}


def test_counter_reset_is_reported_and_skipped(conn, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "Капучино", "Латте"])
    for i, (a, b) in enumerate([(100, 40), (104, 41), (3, 43), (6, 44)]):
        ws.append([date(2025, 9, 1) + timedelta(days=i), a, b])
    path = tmp_path / "счётчики.xlsx"
    wb.save(path)

    data = sales_excel.read_workbook(conn, path)
    daily, notes = data.daily(cumulative=True)
    assert by_name(conn, daily) == {"2025-09-02": {"Капучино": 4, "Латте": 1},
                                    "2025-09-03": {"Латте": 2},
                                    "2025-09-04": {"Капучино": 3, "Латте": 1}}
    assert any("сброшен" in n for n in notes)


def test_transposed_table_and_future_days_skipped(conn, tmp_path):
    """Напитки в строках, дни в столбцах; дни после сегодняшнего (заготовка с формулами) не загружаются."""
    today = date.today()
    days = [today - timedelta(days=2), today - timedelta(days=1), today + timedelta(days=1)]
    wb = Workbook()
    ws = wb.active
    ws.append(["Напиток"] + days)
    ws.append(["Капучино", 3, 4, 4])
    ws.append(["Эспрессо", 1, 0, 0])
    path = tmp_path / "по столбцам.xlsx"
    wb.save(path)

    daily, _ = sales_excel.read_workbook(conn, path).daily(cumulative=False)
    assert by_name(conn, daily) == {days[0].isoformat(): {"Капучино": 3, "Эспрессо": 1},
                                    days[1].isoformat(): {"Капучино": 4}}


def test_import_workbook_writes_sales(conn, tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(["Дата", "Капучино", "Американо"])
    ws.append(["01.09.2025", 2, 1])
    path = tmp_path / "просто.xlsx"
    wb.save(path)

    result = sales_excel.import_workbook(conn, path)
    assert result.days == ["2025-09-01"]
    assert sales_in_db(conn, "2025-09-01") == {"Капучино": 2, "Американо": 1}


def test_file_without_sales_explains_what_is_needed(conn, tmp_path):
    wb = Workbook()
    wb.active.append(["что-то", "совсем", "другое"])
    path = tmp_path / "другое.xlsx"
    wb.save(path)
    data = sales_excel.read_workbook(conn, path)
    assert not data.sheets and data.errors
