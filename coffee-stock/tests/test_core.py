from datetime import date, timedelta

import pytest
from openpyxl import Workbook

from coffee_stock import core, db, sales_excel


@pytest.fixture
def conn(tmp_path):
    return db.connect(tmp_path / "t.db")


def ids(conn, table):
    return {r["name"]: r["id"] for r in conn.execute(f"SELECT id, name FROM {table}")}


def test_seed_has_recipes(conn):
    assert "Капучино" in ids(conn, "drinks")
    assert conn.execute("SELECT COUNT(*) FROM recipe").fetchone()[0] > 0


def test_sales_consume_by_recipe_and_reimport_is_idempotent(conn):
    ing, drk = ids(conn, "ingredients"), ids(conn, "drinks")
    with conn:
        core.add_movement(conn, "2026-10-01", ing["Кофе в зёрнах"], 5000, "initial")
        core.add_movement(conn, "2026-10-01", ing["Молоко сухое"], 2000, "initial")
        core.set_sales_for_date(conn, "2026-10-02", {drk["Капучино"]: 10, drk["Эспрессо"]: 5})
        # Повторная загрузка того же дня с исправленными цифрами
        core.set_sales_for_date(conn, "2026-10-02", {drk["Капучино"]: 20, drk["Эспрессо"]: 5})
    s = core.stock(conn)
    assert s[ing["Кофе в зёрнах"]] == 5000 - 25 * 8
    assert s[ing["Молоко сухое"]] == 2000 - 20 * 10
    assert s[ing["Стаканчики"]] == -25
    assert core.stock(conn, "2026-10-01")[ing["Кофе в зёрнах"]] == 5000


def test_adjust_to_actual(conn):
    ing = ids(conn, "ingredients")["Сахар"]
    with conn:
        core.add_movement(conn, "2026-10-01", ing, 1000, "initial")
        diff = core.adjust_to_actual(conn, "2026-10-05", ing, 950)
    assert diff == -50
    assert core.stock(conn)[ing] == 950


def test_forecast_uses_average_of_recent_days(conn):
    ing, drk = ids(conn, "ingredients"), ids(conn, "drinks")
    start = date.today() - timedelta(days=9)
    with conn:
        core.add_movement(conn, start.isoformat(), ing["Шоколад"], 10000, "initial")
        for i in range(10):
            core.set_sales_for_date(conn, (start + timedelta(days=i)).isoformat(), {drk["Горячий шоколад"]: 4})
    rows = {r.name: r for r in core.forecast(conn, 14)}
    choc = rows["Шоколад"]
    assert choc.per_day == pytest.approx(100)  # 4 порции × 25 г, делим на 10 дней, а не на 14
    assert choc.stock == 10000 - 1000
    assert int(choc.days_left) == 90
    assert rows["Кофе в зёрнах"].days_left is None


def test_portions_left_names_bottleneck(conn):
    ing = ids(conn, "ingredients")
    with conn:
        core.add_movement(conn, "2026-10-01", ing["Шоколад"], 1000, "initial")
        core.add_movement(conn, "2026-10-01", ing["Сахар"], 1000, "initial")
        core.add_movement(conn, "2026-10-01", ing["Стаканчики"], 10, "initial")
        core.add_movement(conn, "2026-10-01", ing["Палочки для размешивания"], 100, "initial")
    left = {r.drink: (r.portions, r.limit) for r in core.portions_left(conn)}
    assert left["Горячий шоколад"] == (10, "Стаканчики")


def test_stirrers_do_not_limit_portions_or_raise_alarm(conn):
    """Палочку берёт примерно каждый второй: норма 0,5, порции и карточку «Заканчивается» она не трогает."""
    ing, drk = ids(conn, "ingredients"), ids(conn, "drinks")
    with conn:
        for name in ("Шоколад", "Сахар"):
            core.add_movement(conn, "2026-10-01", ing[name], 1000, "initial")
        core.add_movement(conn, "2026-10-01", ing["Стаканчики"], 200, "initial")
        core.add_movement(conn, "2026-10-01", ing["Палочки для размешивания"], 2, "initial")
        core.set_sales_for_date(conn, "2026-10-02", {drk["Горячий шоколад"]: 4})
    assert core.stock(conn)[ing["Палочки для размешивания"]] == 0  # 2 − 4 × 0,5
    left = {r.drink: r for r in core.portions_left(conn)}
    assert (left["Горячий шоколад"].portions, left["Горячий шоколад"].limit) == (36, "Шоколад")
    assert left["Горячий шоколад"].per_day == 4
    assert left["Горячий шоколад"].days_left == pytest.approx(900 / 25 / 4)
    rows = {r.name: r for r in core.forecast(conn, 14)}
    assert rows["Палочки для размешивания"].low and rows["Палочки для размешивания"].secondary
    low_main = sum(1 for r in rows.values() if r.low and not r.secondary)
    assert core.summary(conn, 14).low_count == low_main < sum(1 for r in rows.values() if r.low)


def test_reports(conn):
    ing, drk = ids(conn, "ingredients"), ids(conn, "drinks")
    with conn:
        core.add_movement(conn, "2026-10-01", ing["Кофе в зёрнах"], 1000, "initial")
        core.add_movement(conn, "2026-10-03", ing["Кофе в зёрнах"], 500, "receipt", "ручной приход")
        core.set_sales_for_date(conn, "2026-10-02", {drk["Американо"]: 3})
        core.set_sales_for_date(conn, "2026-10-04", {drk["Американо"]: 1, drk["Латте"]: 2})
    header, rows = core.report_sales_by_drink(conn, "2026-10-01", "2026-10-31")
    # Американо 4 × 2,5 = 10, Латте 2 × 3,5 = 7 — сортировка по выручке
    assert rows[0][:3] == ["Американо", 4, 10]
    assert rows[1][:3] == ["Латте", 2, 7]
    assert rows[-1][:3] == ["Итого", 6, 17]
    header, rows = core.report_movement(conn, "2026-10-03", "2026-10-04")
    coffee = next(r for r in rows if r[0] == "Кофе в зёрнах")
    assert coffee[2:] == [1000 - 24, 500, 24, 0, 1000 - 24 + 500 - 24]
    header, rows = core.report_sales_by_day(conn, "2026-10-01", "2026-10-31")
    assert [r[0] for r in rows] == ["2026-10-02", "2026-10-04", "Итого"]
    assert rows[-1][-1] == 17


def test_price_is_fixed_on_the_day_of_sale(conn):
    drk = ids(conn, "drinks")
    with conn:
        core.set_sales_for_date(conn, "2026-10-01", {drk["Капучино"]: 10})
        conn.execute("UPDATE drinks SET price = 4 WHERE id = ?", (drk["Капучино"],))
        # Файл продаж перечитан целиком: старый день остаётся по старой цене
        core.set_sales_for_date(conn, "2026-10-01", {drk["Капучино"]: 10})
        core.set_sales_for_date(conn, "2026-10-02", {drk["Капучино"]: 10})
    money = dict(conn.execute("SELECT date, SUM(count * price) FROM sales GROUP BY date").fetchall())
    assert money == {"2026-10-01": 35, "2026-10-02": 40}


def test_sales_excel_roundtrip(conn, tmp_path):
    path = tmp_path / "Продажи.xlsx"
    sales_excel.ensure_template(conn, path)
    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи"
    ws.append(["Дата", "Капучино", "Эспрессо", "Неизвестный"])
    ws.append([date(2026, 10, 2), 10, None, 1])
    ws.append(["03.10.2026", "4", 2, None])
    ws.append([date(2026, 10, 4), None, None, None])  # ещё не заполнено
    ws.append(["вчера", 1, 1, None])
    wb.save(path)
    result = sales_excel.import_file(conn, path)
    assert result.days == ["2026-10-02", "2026-10-03"]
    assert any("Неизвестный" in e for e in result.errors)
    assert any("вчера" in e for e in result.errors)
    assert conn.execute("SELECT SUM(count) FROM sales").fetchone()[0] == 16


def test_template_adds_new_drink_column(conn, tmp_path):
    path = tmp_path / "Продажи.xlsx"
    sales_excel.ensure_template(conn, path)
    with conn:
        conn.execute("INSERT INTO drinks(name) VALUES ('Раф')")
    sales_excel.ensure_template(conn, path)
    from openpyxl import load_workbook
    header = [c.value for c in load_workbook(path)["Продажи"][1]]
    assert header[0] == "Дата" and header[-1] == "Раф"
