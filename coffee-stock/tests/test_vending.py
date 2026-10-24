"""Кофеавтомат: продажи с автомата → «Продажи.xlsx» и отчёт по кофе (без самого автомата)."""

from datetime import date, datetime, timedelta

from openpyxl import Workbook, load_workbook

from coffee_stock import db, sales_excel, vending


def sale(day: date, product: str, price: float = 2.5, n: int = 0) -> dict:
    return {"order": f"{day:%Y%m%d}-{product}-{n}", "date": day.isoformat(), "time": f"10:{n % 60:02d}:00",
            "product_id": "1", "product": product, "price": price, "pay": "mdb_cashless"}


def sales_file(path, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = sales_excel.SHEET
    ws.append(["Дата", "Эспрессо", "Капучино"])
    for row in rows:
        ws.append(row)
    wb.save(path)


def read(path):
    ws = load_workbook(path).active
    return {(v.date() if isinstance(v, datetime) else v): [c.value for c in row[1:]]
            for row in ws.iter_rows(min_row=2) if (v := row[0].value)}


def test_sync_fills_new_days_updates_last_and_keeps_earlier(tmp_path):
    today = date.today()
    d0, d1, d2, d3 = (today - timedelta(days=i) for i in (3, 2, 1, 0))
    path = tmp_path / "Продажи.xlsx"
    sales_file(path, [[datetime(d0.year, d0.month, d0.day), 50, 40],  # сводная строка — не трогается
                      [datetime(d1.year, d1.month, d1.day), 9, None],  # расходится, но не последняя
                      [datetime(d2.year, d2.month, d2.day), 1, None]])  # последняя — обновится
    month = [sale(d0, "Эспрессо")] + [sale(d1, "Эспрессо", n=i) for i in range(2)] + \
            [sale(d2, "Эспрессо", n=i) for i in range(3)] + [sale(d2, "Капучино", 3.5)] + \
            [sale(d3, "Капучино", 3.5), sale(d3, "Мокко", 4)]
    for ym in sorted({s["date"][:7] for s in month}):
        vending.save_month(tmp_path, ym, "test", [s for s in month if s["date"][:7] == ym])

    changes, notes = vending.sync_sales_file(tmp_path, path)

    got = read(path)
    assert got[d0] == [50, 40]
    assert got[d1] == [9, None]
    assert got[d2] == [3, 1]
    assert got[d3] == [None, 1]
    assert len(changes) == 2
    assert any("мокко" in n for n in notes) and any(f"{d1:%d.%m}" in n for n in notes)
    assert list((tmp_path / "Резерв").iterdir())
    # повторный опрос без новых продаж ничего не меняет
    assert vending.sync_sales_file(tmp_path, path)[0] == []


def test_synced_file_loads_into_program(tmp_path):
    conn = db.connect(tmp_path / "t.db")
    today = date.today()
    path = tmp_path / "Продажи.xlsx"
    sales_excel.ensure_template(conn, path)
    vending.save_month(tmp_path, today.strftime("%Y-%m"), "test",
                       [sale(today, "Капучино", 3.5, i) for i in range(4)])
    vending.sync_sales_file(tmp_path, path)
    result = sales_excel.import_file(conn, path)
    assert today.isoformat() in result.days
    n = conn.execute("SELECT s.count FROM sales s JOIN drinks d ON d.id = s.drink_id "
                     "WHERE d.name = 'Капучино' AND s.date = ?", (today.isoformat(),)).fetchone()[0]
    assert n == 4


def test_summary_and_report(tmp_path):
    day = date(2026, 10, 22)
    vending.save_month(tmp_path, "2026-10", "test",
                       [sale(day, "Капучино", 3.5, i) for i in range(3)] + [sale(day + timedelta(days=1), "Эспрессо")])
    s = vending.summary(tmp_path, "2026-10")
    assert (s.count, s.money) == (4, 13.0)
    assert s.drinks[0][:4] == ("Капучино", "3,50", 3, 10.5)
    assert [d[2] for d in s.days] == [3, 1]
    path = vending.build_report(tmp_path, "2026-10")
    assert path.name == "Кофе 2026-10 октябрь.xlsx"
    assert load_workbook(path).sheetnames == ["По напиткам", "По дням", "Все продажи"]
    assert vending.months(tmp_path) == ["2026-10"]
