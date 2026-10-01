"""Загрузка прайса: Excel (.xlsx/.xls), CSV и выгрузка 1С (CommerceML .xml).

Колонки определяются по заголовкам автоматически; если не угадали —
пользователь выбирает их в окне «Прайс», выбор запоминается.
"""
from __future__ import annotations

import csv
import re
import xml.etree.ElementTree as ET
from pathlib import Path

from .models import PriceItem

ROLES = {
    "name": "Наименование",
    "price": "Цена",
    "code": "Код / артикул",
    "unit": "Ед. изм.",
    "vat": "Ставка НДС",
    "stock": "Остаток",
    "producer": "Производитель",
    "country": "Страна",
}

_ROLE_PATTERNS = {
    "name": [r"наименован", r"номенклатур", r"товар", r"название", r"^продукц"],
    "price": [r"цена", r"стоимость.*ед", r"прайс", r"розн", r"опт"],
    "code": [r"артикул", r"^код", r"код\b", r"арт\.?$"],
    "unit": [r"ед\.?\s*изм", r"^ед\b", r"единиц"],
    "vat": [r"ндс.*%", r"ставка\s*ндс", r"^ндс$"],
    "stock": [r"остат", r"наличи", r"кол-?во на складе", r"склад"],
    "producer": [r"производител", r"изготовител", r"бренд", r"марка"],
    "country": [r"страна"],
}


class PriceTable:
    """Сырая таблица прайса: заголовки + строки (для окна сопоставления колонок)."""

    def __init__(self, headers: list[str], rows: list[list], header_row: int, sheets: list[str], sheet: str):
        self.headers = headers
        self.rows = rows
        self.header_row = header_row
        self.sheets = sheets
        self.sheet = sheet


def _cell_str(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def parse_number(v) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "")
    s = re.sub(r"(руб|br|byn|бел\.?р\.?|р\.)$", "", s, flags=re.I)
    s = s.replace(",", ".")
    if s.count(".") > 1:
        head, _, tail = s.rpartition(".")
        s = head.replace(".", "") + "." + tail
    try:
        return float(s)
    except ValueError:
        m = re.search(r"-?\d+(?:\.\d+)?", s)
        return float(m.group()) if m else None


def _read_raw(path: Path, sheet: str = "") -> tuple[list[list], list[str], str]:
    ext = path.suffix.lower()
    if ext in (".xlsx", ".xlsm"):
        import openpyxl

        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        names = wb.sheetnames
        ws = wb[sheet] if sheet in names else _biggest_sheet(wb)
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        title = ws.title
        wb.close()
        return rows, names, title
    if ext == ".xls":
        import xlrd

        book = xlrd.open_workbook(str(path))
        names = book.sheet_names()
        sh = book.sheet_by_name(sheet) if sheet in names else max(book.sheets(), key=lambda s: s.nrows)
        rows = [sh.row_values(i) for i in range(sh.nrows)]
        return rows, names, sh.name
    if ext in (".csv", ".txt"):
        raw = path.read_bytes()
        for enc in ("utf-8-sig", "cp1251"):
            try:
                text = raw.decode(enc)
                break
            except UnicodeDecodeError:
                continue
        dialect = csv.Sniffer().sniff(text[:5000], delimiters=";,\t")
        rows = [r for r in csv.reader(text.splitlines(), dialect)]
        return rows, [path.name], path.name
    raise ValueError(f"Формат {ext} не поддерживается. Сохраните прайс как .xlsx")


def _biggest_sheet(wb):
    best = wb.worksheets[0]
    for ws in wb.worksheets:
        if (ws.max_row or 0) > (best.max_row or 0):
            best = ws
    return best


def guess_role(header: str) -> str | None:
    h = header.lower().replace("ё", "е").strip()
    if not h:
        return None
    for role in ("vat", "stock", "code", "unit", "producer", "country", "price", "name"):
        for pat in _ROLE_PATTERNS[role]:
            if re.search(pat, h):
                return role
    return None


def detect_header_row(rows: list[list]) -> int:
    best, best_score = 0, -1
    for i, row in enumerate(rows[:40]):
        roles = {guess_role(_cell_str(c)) for c in row} - {None}
        score = len(roles) + (3 if "name" in roles else 0) + (2 if "price" in roles else 0)
        if score > best_score:
            best, best_score = i, score
    return best


def read_table(path: str | Path, sheet: str = "", header_row: int = 0) -> PriceTable:
    path = Path(path)
    rows, sheets, title = _read_raw(path, sheet)
    hr = header_row - 1 if header_row > 0 else detect_header_row(rows)
    headers = [_cell_str(c) or f"Колонка {j + 1}" for j, c in enumerate(rows[hr] if rows else [])]
    # В 1С шапка иногда в две строки: «Цена» над «опт», «розн».
    if hr + 1 < len(rows):
        nxt = [_cell_str(c) for c in rows[hr + 1]]
        if sum(1 for c in nxt if c and not re.search(r"\d", c)) >= 2 and not any(
            parse_number(c) for c in nxt if c
        ):
            headers = [
                (h + " " + n).strip() if n and not h.startswith("Колонка") else (n or h)
                for h, n in zip(headers, nxt + [""] * len(headers))
            ]
            hr += 1
    data = rows[hr + 1 :]
    return PriceTable(headers, data, hr + 1, sheets, title)


def guess_columns(headers: list[str]) -> dict[str, str]:
    result: dict[str, str] = {}
    for h in headers:
        role = guess_role(h)
        if role and role not in result:
            result[role] = h
    return result


def items_from_table(table: PriceTable, columns: dict[str, str]) -> list[PriceItem]:
    idx = {role: table.headers.index(h) for role, h in columns.items() if h in table.headers}
    if "name" not in idx or "price" not in idx:
        raise ValueError("Укажите колонки «Наименование» и «Цена»")

    def get(row, role):
        j = idx.get(role)
        return row[j] if j is not None and j < len(row) else None

    items = []
    for n, row in enumerate(table.rows, start=table.header_row + 1):
        name = _cell_str(get(row, "name"))
        price = parse_number(get(row, "price"))
        if not name or price is None or price <= 0:
            continue  # группы номенклатуры, пустые строки, итоги
        vat = parse_number(get(row, "vat"))
        items.append(
            PriceItem(
                name=name,
                price=round(price, 4),
                code=_cell_str(get(row, "code")),
                unit=_cell_str(get(row, "unit")),
                vat_rate=vat,
                stock=parse_number(get(row, "stock")),
                producer=_cell_str(get(row, "producer")),
                country=_cell_str(get(row, "country")),
                row=n,
            )
        )
    return items


# --- CommerceML (обмен 1С с сайтом: import.xml + offers.xml) --------------------

def _strip_ns(tree):
    for el in tree.iter():
        if isinstance(el.tag, str) and "}" in el.tag:
            el.tag = el.tag.split("}", 1)[1]
    return tree


def load_commerceml(path: Path) -> list[PriceItem]:
    folder = path.parent
    files = [path]
    for other in ("import.xml", "offers.xml"):
        p = folder / other
        if p.exists() and p not in files:
            files.append(p)
    goods: dict[str, dict] = {}
    prices: dict[str, float] = {}
    for f in files:
        root = _strip_ns(ET.parse(f).getroot())
        for t in root.iter("Товар"):
            gid = (t.findtext("Ид") or "").strip()
            goods.setdefault(gid, {}).update(
                name=(t.findtext("Наименование") or "").strip(),
                code=(t.findtext("Артикул") or "").strip(),
                unit=(t.findtext("БазоваяЕдиница") or "").strip(),
            )
        for o in root.iter("Предложение"):
            gid = (o.findtext("Ид") or "").strip()
            g = goods.setdefault(gid, {})
            g.setdefault("name", (o.findtext("Наименование") or "").strip())
            p = o.find(".//Цена/ЦенаЗаЕдиницу")
            if p is not None and p.text:
                prices[gid] = parse_number(p.text) or 0
            q = o.findtext("Количество")
            if q:
                g["stock"] = parse_number(q)
    items = []
    for gid, g in goods.items():
        price = prices.get(gid) or prices.get(gid.split("#")[0])
        if g.get("name") and price:
            items.append(PriceItem(name=g["name"], price=price, code=g.get("code", ""),
                                   unit=g.get("unit", ""), stock=g.get("stock")))
    return items


def load_items(path: str | Path, sheet: str = "", header_row: int = 0,
               columns: dict[str, str] | None = None) -> tuple[list[PriceItem], PriceTable | None]:
    path = Path(path)
    if path.suffix.lower() == ".xml":
        return load_commerceml(path), None
    table = read_table(path, sheet, header_row)
    cols = columns if columns and "name" in columns and "price" in columns else guess_columns(table.headers)
    cols = {r: h for r, h in cols.items() if h in table.headers}
    return items_from_table(table, cols), table
