"""Заполнение форм заказчика (.docx/.xlsx): таблица ценового предложения
и пустые поля с реквизитами («Наименование участника: ________»).
"""
from __future__ import annotations

import copy
import re
from datetime import date
from pathlib import Path

from rapidfuzz import fuzz

from .. import textnorm
from ..config import OfferTerms, Requisites
from ..models import Match

# --- форматирование -----------------------------------------------------------


def money(v: float | None) -> str:
    if v is None:
        return ""
    s = f"{v:,.2f}".replace(",", " ").replace(".", ",")
    return s


def qty_str(v: float | None) -> str:
    if v is None:
        return ""
    return str(int(v)) if float(v).is_integer() else f"{v:g}".replace(".", ",")


# --- расчёт строк -------------------------------------------------------------


class OfferLine:
    def __init__(self, n: int, m: Match, terms: OfferTerms, req: Requisites):
        self.n = n
        self.m = m
        self.name = m.position.name
        self.offered = m.item.name if m.item else ""
        self.unit = m.position.unit or (m.item.unit if m.item else "")
        self.qty = m.qty
        self.price = round(m.price or 0, 2)
        self.sum = round(self.price * self.qty, 2)
        rate = 0.0
        if req.vat_payer:
            rate = m.item.vat_rate if m.item and m.item.vat_rate is not None else terms.vat_rate
        self.vat_rate = rate
        self.vat_sum = round(self.sum * rate / 100, 2)
        self.total = round(self.sum + self.vat_sum, 2)
        self.producer = (m.item.producer if m.item else "") or ""
        self.country = (m.item.country if m.item else "") or req.country_of_origin
        self.term = terms.delivery_term
        self.code = m.item.code if m.item else ""

    def value(self, role: str, req: Requisites) -> str:
        return {
            "num": str(self.n),
            "name": self.name,
            "offered": self.offered,
            "unit": self.unit,
            "qty": qty_str(self.qty),
            "price": money(self.price),
            "sum": money(self.sum),
            "vat_rate": (f"{self.vat_rate:g}%" if req.vat_payer else "без НДС"),
            "vat_sum": (money(self.vat_sum) if req.vat_payer else "—"),
            "total": money(self.total),
            "price_vat": money(round(self.price * (1 + self.vat_rate / 100), 2)),
            "producer": self.producer,
            "country": self.country,
            "term": self.term,
            "code": self.code,
        }.get(role, "")


def offer_lines(matches: list[Match], terms: OfferTerms, req: Requisites) -> list[OfferLine]:
    chosen = [m for m in matches if m.found and m.include]
    return [OfferLine(i + 1, m, terms, req) for i, m in enumerate(chosen)]


def totals(lines: list[OfferLine]) -> dict[str, float]:
    return {
        "sum": round(sum(x.sum for x in lines), 2),
        "vat_sum": round(sum(x.vat_sum for x in lines), 2),
        "total": round(sum(x.total for x in lines), 2),
    }


# --- распознавание колонок формы ---------------------------------------------

_ROLES = [
    ("num", r"^(№|n|номер)\s*(п/?п|лота)?\.?$|^№\s*п|^п/п"),
    ("vat_rate", r"ставк\w*\s*ндс|ндс\s*,?\s*%"),
    ("vat_sum", r"сумм\w*\s*ндс|ндс\s*,?\s*(руб|byn|бел)"),
    ("total", r"(с\s*(учетом\s*)?ндс|всего|итого)"),
    ("price_vat", r"цен\w*.*с\s*ндс"),
    ("price", r"цен\w*"),
    ("sum", r"(стоимост|сумм)\w*"),
    ("qty", r"кол-?во|количеств|объ[её]м"),
    ("unit", r"ед\.?\s*изм|^ед\b|единиц"),
    ("country", r"страна"),
    ("producer", r"производител|изготовител"),
    ("term", r"срок"),
    ("offered", r"(предлагаем|марка|модель|характеристик\w* предлаг|торгов\w* (марк|наимен))"),
    ("code", r"артикул|код\s*товара"),
    ("name", r"наимен|предмет|товар|номенклатур|описание"),
]


def classify_header(h: str) -> str | None:
    hl = re.sub(r"\s+", " ", h.lower().replace("ё", "е")).strip()
    if not hl:
        return None
    if re.search(r"цен\w*.*с\s*ндс", hl):
        return "price_vat"
    if re.search(r"(стоимост|сумм)\w*.*с\s*(учетом\s*)?ндс", hl):
        return "total"
    if re.search(r"(стоимост|сумм|цен)\w*.*без\s*ндс", hl):
        return "price" if "цен" in hl else "sum"
    for role, pat in _ROLES:
        if re.search(pat, hl):
            return role
    return None


def map_offer_header(cells: list[str]) -> dict[int, str]:
    roles: dict[int, str] = {}
    used = set()
    seen_text = None
    for j, c in enumerate(cells):
        if c == seen_text:
            continue  # объединённые ячейки повторяются
        seen_text = c
        r = classify_header(c)
        if r and r not in used:
            roles[j] = r
            used.add(r)
    if "name" not in used and "offered" in used:
        for j, r in list(roles.items()):
            if r == "offered":
                roles[j] = "name"
    return roles


def _is_offer_header(roles: dict[int, str]) -> bool:
    vals = set(roles.values())
    return "name" in vals and bool(vals & {"price", "sum", "total", "price_vat"})


def _is_numbering_row(cells: list[str]) -> bool:
    vals = [c.strip() for c in cells if c.strip()]
    return bool(vals) and all(re.fullmatch(r"\d{1,2}", v) for v in vals)


def _row_line(name: str, lines: list[OfferLine], used: set[int]) -> OfferLine | None:
    best, best_s = None, 0
    k = textnorm.key(name)
    for i, ln in enumerate(lines):
        if i in used:
            continue
        s = fuzz.token_set_ratio(k, textnorm.key(ln.name))
        if s > best_s:
            best, best_s = i, s
    if best is not None and best_s >= 80:
        used.add(best)
        return lines[best]
    return None


# --- реквизиты в тексте -------------------------------------------------------

def requisite_values(req: Requisites, terms: OfferTerms) -> list[tuple[str, str]]:
    bank = ", ".join(x for x in [
        f"р/с {req.bank_account}" if req.bank_account else "",
        req.bank_name, f"BIC {req.bank_bic}" if req.bank_bic else "", req.bank_address] if x)
    director = " ".join(x for x in [req.director_position, req.director_name] if x)
    # Порядок важен: сначала более конкретные подписи.
    return [
        (r"сокращенн\w* наименовани", req.short_name or req.full_name),
        (r"(полное )?наименовани\w* (участника|организации|поставщика|претендента|юридического лица)"
         r"|^наименовани\w*$|фирменн\w* наименовани", req.full_name),
        (r"(учетн\w* номер плательщика|унп)", req.unp),
        (r"окпо", req.okpo),
        (r"(юридическ\w* адрес|место нахождени|местонахождени)", req.legal_address),
        (r"почтов\w* адрес", req.postal_address or req.legal_address),
        (r"(банковск\w* реквизит|платежн\w* реквизит)", bank),
        (r"(расчетн\w* сч|р/с|iban|текущ\w* \(?расчетн)", req.bank_account),
        (r"\b(bic|бик)\b", req.bank_bic),
        (r"(наименовани\w* банка|^банк\b|обслуживающ\w* банк)", req.bank_name),
        (r"(e-?mail|электронн\w* почт|эл\. ?почт)", req.email),
        (r"(факс|телефон|тел\.)", req.phone),
        (r"(контактн\w* лиц)", ", ".join(x for x in [req.contact_person, req.phone] if x)),
        (r"(руководител|ф\.?и\.?о\.? руководител|должность.*руководител)", director),
        (r"(свидетельств\w* о (гос\w* )?регистрации|сведения о регистрации)", req.registration_info),
        (r"(сайт|web)", req.website),
        (r"срок\w* поставки", terms.delivery_term),
        (r"(условия|порядок) оплаты", terms.payment_terms),
        (r"срок\w* действия (предложения|ценового|конкурсного)", f"{terms.offer_validity_days} календарных дней"),
        (r"(гаранти\w* срок|срок\w* гаранти)", terms.warranty),
        (r"(условия|место|базис) поставки", terms.delivery_terms_place),
        (r"страна происхождени", req.country_of_origin),
        (r"(статус участника|участник является)", ("производитель" if req.is_producer else "поставщик")),
    ]


def value_for_label(label: str, pairs: list[tuple[str, str]]) -> str:
    l = re.sub(r"\s+", " ", label.lower().replace("ё", "е")).strip(" :*")
    for pat, val in pairs:
        if val and re.search(pat, l):
            return val
    return ""


_BLANK = re.compile(r"_{3,}|\.{6,}|…{2,}")


def _set_paragraph_text(p, text: str) -> None:
    runs = p.runs
    if not runs:
        p.add_run(text)
        return
    runs[0].text = text
    for r in runs[1:]:
        r.text = ""


def _fill_paragraph_blanks(p, pairs) -> int:
    text = p.text
    if not _BLANK.search(text):
        return 0
    m = re.match(r"\s*([^_….]{3,120}?)\s*[:\-–]?\s*(_{3,}|\.{6,}|…{2,})", text)
    if not m:
        return 0
    val = value_for_label(m.group(1), pairs)
    if not val:
        return 0
    new = text[: m.start(2)] + val + text[m.end(2):]
    _set_paragraph_text(p, new)
    return 1


def _set_cell(cell, text: str) -> None:
    paras = cell.paragraphs
    if not paras:
        cell.text = text
        return
    _set_paragraph_text(paras[0], text)
    for extra in paras[1:]:
        _set_paragraph_text(extra, "")


# --- DOCX ---------------------------------------------------------------------

def fill_docx(src: Path, dst: Path, lines: list[OfferLine], req: Requisites, terms: OfferTerms,
              log=print) -> dict:
    import docx

    d = docx.Document(str(src))
    pairs = requisite_values(req, terms)
    stats = {"rows": 0, "fields": 0, "offer_table": False}

    for p in d.paragraphs:
        stats["fields"] += _fill_paragraph_blanks(p, pairs)

    for table in d.tables:
        rows = table.rows
        header_i, roles = None, {}
        for i, r in enumerate(rows[:5]):
            try:
                cells = [c.text.strip() for c in r.cells]
            except Exception:
                break
            rr = map_offer_header(cells)
            if _is_offer_header(rr):
                header_i, roles = i, rr
                # Подзаголовки в две строки: «Стоимость» → «без НДС | НДС | с НДС»
                if i + 1 < len(rows):
                    sub = [c.text.strip() for c in rows[i + 1].cells]
                    sub_roles = map_offer_header(sub)
                    if len(set(sub_roles.values()) - {"name"}) >= 2 and not _is_numbering_row(sub):
                        for j, role in sub_roles.items():
                            roles[j] = role
                        header_i = i + 1
                break
        if header_i is None:
            for r in rows:  # анкетная таблица: «подпись | пусто»
                try:
                    cells = r.cells
                except Exception:
                    continue
                if len(cells) >= 2 and len(cells[0].text) < 150 and not cells[-1].text.strip():
                    val = value_for_label(cells[0].text, pairs)
                    if val:
                        _set_cell(cells[-1], val)
                        stats["fields"] += 1
            continue

        stats["offer_table"] = True
        body = list(rows[header_i + 1:])
        if body and _is_numbering_row([c.text for c in body[0].cells]):
            body = body[1:]
        total_rows = [r for r in body if re.match(r"\s*(итого|всего)", " ".join(c.text for c in r.cells[:3]).lower())]
        data_rows = [r for r in body if r not in total_rows]
        name_col = next(j for j, r in roles.items() if r == "name")
        prefilled = [r for r in data_rows if len(r.cells) > name_col and len(r.cells[name_col].text.strip()) > 3]

        if prefilled:
            used: set[int] = set()
            for r in prefilled:
                ln = _row_line(r.cells[name_col].text, lines, used)
                if not ln:
                    continue
                for j, role in roles.items():
                    if role in ("name", "num") or j >= len(r.cells):
                        continue
                    if role in ("qty", "unit") and r.cells[j].text.strip():
                        continue  # заказчик уже указал
                    _set_cell(r.cells[j], ln.value(role, req))
                stats["rows"] += 1
        else:
            empty = [r for r in data_rows if not "".join(c.text for c in r.cells).strip()
                     or _is_numbering_row([c.text for c in r.cells])]
            template = empty[0] if empty else (data_rows[-1] if data_rows else rows[header_i])
            anchor = total_rows[0]._tr if total_rows else None
            for k, ln in enumerate(lines):
                if k < len(empty):
                    r = empty[k]
                else:
                    new_tr = copy.deepcopy(template._tr)
                    if anchor is not None:
                        anchor.addprevious(new_tr)
                    else:
                        table._tbl.append(new_tr)
                    r = _row_for_tr(table, new_tr)
                for j, role in roles.items():
                    if j < len(r.cells):
                        _set_cell(r.cells[j], ln.value(role, req))
                stats["rows"] += 1
            for r in empty[len(lines):]:
                r._tr.getparent().remove(r._tr)

        t = totals(lines)
        for r in total_rows:
            for j, role in roles.items():
                if j >= len(r.cells):
                    continue
                key = {"sum": "sum", "vat_sum": "vat_sum", "total": "total"}.get(role)
                if key and r.cells[j].text.strip() in ("", "-", "—"):
                    _set_cell(r.cells[j], money(t[key]))

    _replace_placeholders_docx(d, req, terms, lines)
    dst.parent.mkdir(parents=True, exist_ok=True)
    d.save(str(dst))
    return stats


def _row_for_tr(table, tr):
    for r in table.rows:
        if r._tr is tr:
            return r
    raise RuntimeError("строка не найдена")


def _replace_placeholders_docx(d, req: Requisites, terms: OfferTerms, lines) -> None:
    t = totals(lines)
    repl = {
        "«___»": f"«{date.today().day:02d}»",
        "{{ДАТА}}": date.today().strftime("%d.%m.%Y"),
        "{{ИТОГО}}": money(t["total"]),
    }
    for p in d.paragraphs:
        txt = p.text
        if any(k in txt for k in repl):
            for k, v in repl.items():
                txt = txt.replace(k, v)
            _set_paragraph_text(p, txt)


# --- XLSX ---------------------------------------------------------------------

def fill_xlsx(src: Path, dst: Path, lines: list[OfferLine], req: Requisites, terms: OfferTerms,
              log=print) -> dict:
    import openpyxl
    from openpyxl.cell.cell import MergedCell

    wb = openpyxl.load_workbook(src)
    pairs = requisite_values(req, terms)
    stats = {"rows": 0, "fields": 0, "offer_table": False}

    def put(ws, r, c, val):
        cell = ws.cell(r, c)
        if isinstance(cell, MergedCell):
            return
        if isinstance(cell.value, str) and cell.value.startswith("="):
            return  # формула заказчика посчитает сама
        num = _to_number(val)
        cell.value = num if num is not None else val

    for ws in wb.worksheets:
        header_r, roles = None, {}
        for r in range(1, min(ws.max_row, 40) + 1):
            cells = [str(ws.cell(r, c).value or "").strip() for c in range(1, ws.max_column + 1)]
            rr = map_offer_header(cells)
            if _is_offer_header(rr):
                header_r, roles = r, rr
                if r + 1 <= ws.max_row:
                    sub = [str(ws.cell(r + 1, c).value or "").strip() for c in range(1, ws.max_column + 1)]
                    sr = map_offer_header(sub)
                    if len(set(sr.values()) - {"name"}) >= 2 and not _is_numbering_row(sub):
                        roles.update(sr)
                        header_r = r + 1
                break
        if header_r is None:
            for r in range(1, min(ws.max_row, 200) + 1):
                for c in range(1, min(ws.max_column, 10)):
                    lab = ws.cell(r, c).value
                    if isinstance(lab, str) and 2 < len(lab) < 150:
                        target = ws.cell(r, c + 1)
                        if not isinstance(target, MergedCell) and target.value in (None, ""):
                            val = value_for_label(lab, pairs)
                            if val:
                                target.value = val
                                stats["fields"] += 1
            continue

        stats["offer_table"] = True
        name_c = next(j for j, r in roles.items() if r == "name") + 1
        start = header_r + 1
        first = [str(ws.cell(start, c).value or "") for c in range(1, ws.max_column + 1)]
        if _is_numbering_row(first):
            start += 1
        end = start
        total_row = None
        while end <= ws.max_row:
            txt = " ".join(str(ws.cell(end, c).value or "") for c in range(1, 4)).lower()
            if re.match(r"\s*(итого|всего)", txt.strip()):
                total_row = end
                break
            end += 1
        existing = [r for r in range(start, end) if len(str(ws.cell(r, name_c).value or "").strip()) > 3]
        if existing:
            used: set[int] = set()
            for r in existing:
                ln = _row_line(str(ws.cell(r, name_c).value), lines, used)
                if not ln:
                    continue
                for j, role in roles.items():
                    if role in ("name", "num"):
                        continue
                    if role in ("qty", "unit") and ws.cell(r, j + 1).value not in (None, ""):
                        continue
                    put(ws, r, j + 1, ln.value(role, req))
                stats["rows"] += 1
        else:
            free = end - start
            if total_row and len(lines) > free:
                ws.insert_rows(total_row, len(lines) - free)
                total_row += len(lines) - free
            for k, ln in enumerate(lines):
                for j, role in roles.items():
                    put(ws, start + k, j + 1, ln.value(role, req))
                stats["rows"] += 1
        if total_row:
            t = totals(lines)
            for j, role in roles.items():
                if role in t:
                    put(ws, total_row, j + 1, t[role])

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return stats


def _to_number(val: str):
    if not isinstance(val, str):
        return val
    s = val.replace(" ", "").replace("\xa0", "")
    if re.fullmatch(r"-?\d+(,\d+)?", s):
        return float(s.replace(",", "."))
    return None


def fill_form(src: Path, dst: Path, lines, req, terms, log=print) -> dict:
    if src.suffix.lower() == ".docx":
        return fill_docx(src, dst, lines, req, terms, log)
    if src.suffix.lower() in (".xlsx", ".xlsm"):
        return fill_xlsx(src, dst, lines, req, terms, log)
    raise ValueError(f"Не умею заполнять {src.suffix}")
