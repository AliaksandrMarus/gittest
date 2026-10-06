"""Документы, которые программа составляет сама по нашим реквизитам:
ценовое предложение (если у заказчика нет своей формы или как запасной
вариант), сопроводительное письмо, сведения об участнике, опись,
сравнение с прайсом и чек-лист перед подачей.
"""
from __future__ import annotations

from datetime import date, timedelta
from pathlib import Path

from docx import Document as Docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

from ..config import OfferTerms, Requisites
from ..models import Match, Tender
from .fill import OfferLine, money, qty_str, totals

# --- общие элементы -----------------------------------------------------------


def _doc() -> Docx:
    d = Docx()
    st = d.styles["Normal"]
    st.font.name = "Times New Roman"
    st.font.size = Pt(12)
    rpr = st.element.get_or_add_rPr()
    rfonts = rpr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}rFonts")
    if rfonts is not None:
        rfonts.set("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia", "Times New Roman")
    for s in d.sections:
        s.left_margin, s.right_margin = Cm(2.5), Cm(1.5)
        s.top_margin, s.bottom_margin = Cm(2), Cm(2)
    return d


def _letterhead(d: Docx, req: Requisites) -> None:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(req.full_name or "НАИМЕНОВАНИЕ ОРГАНИЗАЦИИ (заполните реквизиты)")
    r.bold = True
    info = ", ".join(x for x in [req.legal_address, f"УНП {req.unp}" if req.unp else "",
                                  f"тел. {req.phone}" if req.phone else "", req.email] if x)
    if info:
        p2 = d.add_paragraph(info)
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p2.runs[0].font.size = Pt(10)
    bank = ", ".join(x for x in [f"р/с {req.bank_account}" if req.bank_account else "", req.bank_name,
                                  f"BIC {req.bank_bic}" if req.bank_bic else ""] if x)
    if bank:
        p3 = d.add_paragraph(bank)
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p3.runs[0].font.size = Pt(10)
    d.add_paragraph("_" * 75).alignment = WD_ALIGN_PARAGRAPH.CENTER


def _signature(d: Docx, req: Requisites) -> None:
    d.add_paragraph()
    t = d.add_table(rows=1, cols=3)
    c = t.rows[0].cells
    c[0].text = req.director_position or "Руководитель"
    c[1].text = "______________"
    c[2].text = req.director_short or req.director_name or "_____________"
    d.add_paragraph("М.П. (при наличии)")
    d.add_paragraph(f"«{date.today().day:02d}» {_month(date.today())} {date.today().year} г.")


_MONTHS = ["января", "февраля", "марта", "апреля", "мая", "июня", "июля", "августа",
           "сентября", "октября", "ноября", "декабря"]


def _month(d: date) -> str:
    return _MONTHS[d.month - 1]


def _heading(d: Docx, text: str) -> None:
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(14)


def _grid(table) -> None:
    table.style = "Table Grid"


# --- ценовое предложение ------------------------------------------------------


def offer_docx(path: Path, t: Tender, lines: list[OfferLine], req: Requisites, terms: OfferTerms) -> None:
    d = _doc()
    _letterhead(d, req)
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(f"{t.customer or 'Заказчику'}")
    _heading(d, "ЦЕНОВОЕ ПРЕДЛОЖЕНИЕ")
    d.add_paragraph(
        f"Изучив документацию по процедуре закупки № {t.number or t.ext_id} «{t.title}», "
        f"{req.full_name or '[наименование участника]'} предлагает поставить товар на следующих условиях:"
    )
    vat = req.vat_payer
    headers = ["№", "Наименование товара", "Ед. изм.", "Кол-во", "Цена за ед. без НДС, руб.",
               "Стоимость без НДС, руб."]
    if vat:
        headers += ["Ставка НДС", "Сумма НДС, руб.", "Стоимость с НДС, руб."]
    headers += ["Страна происхождения"]
    tb = d.add_table(rows=1, cols=len(headers))
    _grid(tb)
    for j, h in enumerate(headers):
        tb.rows[0].cells[j].text = h
        for r in tb.rows[0].cells[j].paragraphs[0].runs:
            r.bold = True
            r.font.size = Pt(9)
    for ln in lines:
        name = ln.name if not ln.offered or ln.offered == ln.name else f"{ln.name}\n(предлагается: {ln.offered})"
        vals = [str(ln.n), name, ln.unit, qty_str(ln.qty), money(ln.price), money(ln.sum)]
        if vat:
            vals += [f"{ln.vat_rate:g}%", money(ln.vat_sum), money(ln.total)]
        vals += [ln.country]
        row = tb.add_row().cells
        for j, v in enumerate(vals):
            row[j].text = v
            for r in row[j].paragraphs[0].runs:
                r.font.size = Pt(9)
    tt = totals(lines)
    row = tb.add_row().cells
    row[1].text = "ИТОГО"
    row[5].text = money(tt["sum"])
    if vat:
        row[7].text = money(tt["vat_sum"])
        row[8].text = money(tt["total"])
    d.add_paragraph()
    total_text = money(tt["total"] if vat else tt["sum"])
    d.add_paragraph(f"Общая стоимость предложения: {total_text} бел. руб."
                    + (f", в том числе НДС {money(tt['vat_sum'])} бел. руб." if vat else " (без НДС)."))
    for label, val in [
        ("Срок поставки", terms.delivery_term),
        ("Условия поставки", terms.delivery_terms_place),
        ("Условия оплаты", terms.payment_terms),
        ("Гарантийный срок", terms.warranty),
        ("Срок действия предложения",
         f"{terms.offer_validity_days} календарных дней (до {(date.today() + timedelta(days=terms.offer_validity_days)).strftime('%d.%m.%Y')})"),
        ("Валюта предложения", "белорусский рубль"),
    ]:
        d.add_paragraph(f"{label}: {val}.")
    d.add_paragraph("Цена включает все расходы участника, связанные с исполнением договора, "
                    "в том числе налоги, сборы, расходы на доставку и разгрузку.")
    d.add_paragraph("Настоящим подтверждаем, что предлагаемый товар соответствует требованиям "
                    "документации о закупке.")
    _signature(d, req)
    d.save(str(path))


def offer_xlsx(path: Path, t: Tender, lines: list[OfferLine], req: Requisites) -> None:
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, Side

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Предложение"
    ws["A1"] = f"Ценовое предложение {req.short_name or req.full_name} по процедуре № {t.number or t.ext_id}"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = t.title
    vat = req.vat_payer
    headers = ["№", "Наименование (заказчик)", "Предлагаемый товар", "Код", "Ед.", "Кол-во",
               "Цена без НДС", "Стоимость без НДС"]
    if vat:
        headers += ["НДС %", "Сумма НДС", "Стоимость с НДС"]
    headers += ["Страна"]
    ws.append([])
    ws.append(headers)
    thin = Side(style="thin")
    for c in ws[4]:
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for ln in lines:
        row = [ln.n, ln.name, ln.offered, ln.code, ln.unit, ln.qty, ln.price, ln.sum]
        if vat:
            row += [ln.vat_rate, ln.vat_sum, ln.total]
        row += [ln.country]
        ws.append(row)
    tt = totals(lines)
    total_row = ["", "ИТОГО", "", "", "", "", "", tt["sum"]]
    if vat:
        total_row += ["", tt["vat_sum"], tt["total"]]
    ws.append(total_row)
    for row in ws.iter_rows(min_row=4, max_row=ws.max_row):
        for c in row:
            c.border = Border(top=thin, bottom=thin, left=thin, right=thin)
            if isinstance(c.value, float):
                c.number_format = "#,##0.00"
    for col, w in zip("ABCDEFGHIJKL", [5, 45, 45, 12, 7, 9, 13, 15, 7, 13, 15, 18]):
        ws.column_dimensions[col].width = w
    wb.save(path)


# --- сопроводительные документы ----------------------------------------------


def cover_letter(path: Path, t: Tender, req: Requisites, attachments: list[str], total: float) -> None:
    d = _doc()
    _letterhead(d, req)
    d.add_paragraph(f"Исх. № ____ от {date.today().strftime('%d.%m.%Y')}")
    p = d.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.add_run(t.customer or "Заказчику")
    _heading(d, "Сопроводительное письмо")
    d.add_paragraph(
        f"{req.full_name or '[наименование участника]'} направляет предложение для участия "
        f"в процедуре закупки № {t.number or t.ext_id} «{t.title}» на общую сумму "
        f"{money(total)} бел. руб.{' с НДС' if req.vat_payer else ' (без НДС)'}."
    )
    d.add_paragraph("Приложения:")
    for i, a in enumerate(attachments, 1):
        d.add_paragraph(f"{i}. {a}")
    _signature(d, req)
    d.save(str(path))


def participant_info(path: Path, req: Requisites, terms: OfferTerms) -> None:
    d = _doc()
    _heading(d, "СВЕДЕНИЯ ОБ УЧАСТНИКЕ")
    rows = [
        ("Полное наименование", req.full_name),
        ("Сокращённое наименование", req.short_name),
        ("УНП", req.unp),
        ("ОКПО", req.okpo),
        ("Сведения о государственной регистрации", req.registration_info),
        ("Юридический адрес", req.legal_address),
        ("Почтовый адрес", req.postal_address or req.legal_address),
        ("Банковские реквизиты", ", ".join(x for x in [req.bank_account, req.bank_name,
                                                        f"BIC {req.bank_bic}" if req.bank_bic else "",
                                                        req.bank_address] if x)),
        ("Руководитель", " ".join(x for x in [req.director_position, req.director_name] if x)),
        ("Действует на основании", req.acts_on),
        ("Контактное лицо", req.contact_person),
        ("Телефон", req.phone),
        ("Электронная почта", req.email),
        ("Сайт", req.website),
        ("Плательщик НДС", "да" if req.vat_payer else "нет"),
        ("Резидент Республики Беларусь", "да" if req.is_resident else "нет"),
        ("Статус", "производитель" if req.is_producer else "поставщик (не производитель)"),
    ]
    tb = d.add_table(rows=0, cols=2)
    _grid(tb)
    for k, v in rows:
        c = tb.add_row().cells
        c[0].text = k
        c[1].text = v or ""
    _signature(d, req)
    d.save(str(path))


def inventory(path: Path, t: Tender, req: Requisites, items: list[str]) -> None:
    d = _doc()
    _heading(d, "ОПИСЬ ДОКУМЕНТОВ")
    d.add_paragraph(f"представленных {req.short_name or req.full_name or '[участник]'} для участия "
                    f"в процедуре закупки № {t.number or t.ext_id}")
    tb = d.add_table(rows=1, cols=3)
    _grid(tb)
    for j, h in enumerate(["№", "Наименование документа", "Кол-во листов"]):
        tb.rows[0].cells[j].text = h
    for i, name in enumerate(items, 1):
        c = tb.add_row().cells
        c[0].text, c[1].text, c[2].text = str(i), name, ""
    _signature(d, req)
    d.save(str(path))


def comparison_xlsx(path: Path, matches: list[Match]) -> None:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Сравнение"
    ws.append(["№", "Позиция заказчика", "Кол-во", "Ед.", "Источник", "Найдено в прайсе", "Код",
               "Похожесть %", "Цена прайса", "Цена в предложении", "Включено", "Ориентир. стоимость"])
    for c in ws[1]:
        c.font = Font(bold=True)
    red = PatternFill("solid", fgColor="F8D7DA")
    yellow = PatternFill("solid", fgColor="FFF3CD")
    for i, m in enumerate(matches, 1):
        ws.append([i, m.position.name, m.position.qty, m.position.unit, m.position.source,
                   m.item.name if m.item else "НЕ НАЙДЕНО", m.item.code if m.item else "",
                   m.score, m.item.price if m.item else None, m.price,
                   "да" if (m.found and m.include) else "нет", m.position.price_limit])
        fill = red if not m.found else (yellow if m.score < 85 and not m.manual else None)
        if fill:
            for c in ws[ws.max_row]:
                c.fill = fill
    for col, w in zip("ABCDEFGHIJKL", [5, 50, 9, 7, 18, 50, 12, 11, 12, 14, 9, 14]):
        ws.column_dimensions[col].width = w
    wb.save(path)


def checklist(path: Path, t: Tender, matches: list[Match], required: list[str],
              conditions: dict[str, str], filled: list[str], notes: list[str]) -> None:
    d = _doc()
    _heading(d, "ЧЕК-ЛИСТ ПЕРЕД ПОДАЧЕЙ")
    d.add_paragraph(f"Процедура № {t.number or t.ext_id}: {t.title}")
    d.add_paragraph(f"Заказчик: {t.customer}")
    if t.deadline:
        d.add_paragraph(f"Окончание приёма предложений: {t.deadline.replace('T', ' ')}")
    d.add_paragraph(f"Ссылка: {t.url}")
    d.add_paragraph()
    d.add_paragraph("Подготовлено программой:").runs[0].bold = True
    for f in filled:
        d.add_paragraph(f"☐ {f} — проверить и подписать ЭЦП")
    missing = [m for m in matches if not m.found]
    weak = [m for m in matches if m.found and m.score < 85 and not m.manual]
    if missing or weak:
        d.add_paragraph()
        d.add_paragraph("Проверьте позиции:").runs[0].bold = True
        for m in missing:
            d.add_paragraph(f"☐ НЕТ В ПРАЙСЕ: {m.position.name}")
        for m in weak:
            d.add_paragraph(f"☐ Сомнительное совпадение ({m.score}%): «{m.position.name}» → «{m.item.name}»")
    if required:
        d.add_paragraph()
        d.add_paragraph("Документация упоминает (приложить, если требуется):").runs[0].bold = True
        for r in required:
            d.add_paragraph(f"☐ {r}")
    if conditions:
        d.add_paragraph()
        d.add_paragraph("Условия из документации — сверьте с предложением:").runs[0].bold = True
        for k, v in conditions.items():
            d.add_paragraph(f"• {k}: {v}")
    if notes:
        d.add_paragraph()
        d.add_paragraph("Замечания программы:").runs[0].bold = True
        for n in notes:
            d.add_paragraph(f"• {n}")
    d.add_paragraph()
    d.add_paragraph("Подача: откройте страницу процедуры на площадке, загрузите файлы из папки "
                    "«2. Наше предложение (для подачи)», подпишите ЭЦП (ключ Авест) и отправьте.")
    d.save(str(path))
