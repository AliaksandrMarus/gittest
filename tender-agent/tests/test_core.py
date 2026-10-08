import os
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))


@pytest.fixture(autouse=True)
def home(tmp_path, monkeypatch):
    monkeypatch.setenv("TENDERAGENT_HOME", str(tmp_path / "home"))
    yield tmp_path


def make_price(path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Прайс-лист ООО «Ромашка» на 01.10.2026"])
    ws.append([])
    ws.append(["Код", "Номенклатура", "Ед. изм.", "Цена, руб.", "Остаток"])
    ws.append(["", "Канцтовары", "", "", ""])  # группа 1С без цены
    ws.append(["1001", "Бумага офисная А4 80 г/м2 500 л. SvetoCopy", "пачка", "12,50", 100])
    ws.append(["1002", "Ручка шариковая синяя BIC", "шт", 0.85, 500])
    ws.append(["1003", "Кабель ВВГнг 3х2,5", "м", 3.2, 1000])
    ws.append(["1004", "Кабель ВВГнг 3х1,5", "м", 2.1, 1000])
    ws.append(["1005", "Степлер №24/6 металлический", "шт", 9.9, 10])
    wb.save(path)


def test_pricelist_autodetect(tmp_path):
    from tenderagent import pricelist

    p = tmp_path / "price.xlsx"
    make_price(p)
    items, table = pricelist.load_items(p)
    assert len(items) == 5
    assert items[0].price == 12.5 and items[0].code == "1001" and items[0].unit == "пачка"


def test_matching_prefers_numbers(tmp_path):
    from tenderagent import pricelist
    from tenderagent.config import OfferTerms
    from tenderagent.matching import PriceIndex, match_positions
    from tenderagent.models import Position

    p = tmp_path / "price.xlsx"
    make_price(p)
    items, _ = pricelist.load_items(p)
    idx = PriceIndex(items)
    pos = [Position("Кабель силовой ВВГнг 3x1.5", qty=200, unit="м"),
           Position("Бумага для офисной техники формат А4, 80 г/м2", qty=50),
           Position("Ноутбук игровой", qty=1)]
    ms = match_positions(pos, idx, OfferTerms(), 60)
    assert ms[0].item.code == "1004"
    assert ms[1].item.code == "1001"
    assert ms[2].item is None


def make_offer_form(path: Path):
    import docx

    d = docx.Document()
    d.add_paragraph("Приложение 2. Форма ценового предложения")
    d.add_paragraph("Наименование участника: ____________________")
    d.add_paragraph("УНП: __________")
    t = d.add_table(rows=1, cols=7)
    for j, h in enumerate(["№ п/п", "Наименование товара", "Ед. изм.", "Количество",
                           "Цена за единицу без НДС, руб.", "Стоимость без НДС, руб.", "Стоимость с НДС, руб."]):
        t.rows[0].cells[j].text = h
    for i, (n, q) in enumerate([("Бумага офисная А4", "50"), ("Ручка шариковая синяя", "100"),
                                ("Маркер текстовыделитель", "10")], 1):
        r = t.add_row().cells
        r[0].text, r[1].text, r[2].text, r[3].text = str(i), n, "шт", q
    r = t.add_row().cells
    r[1].text = "Итого"
    d.add_paragraph("Анкета")
    a = d.add_table(rows=2, cols=2)
    a.rows[0].cells[0].text = "Юридический адрес"
    a.rows[1].cells[0].text = "Банковские реквизиты"
    d.save(path)


def test_fill_customer_form(tmp_path):
    import docx

    from tenderagent import pricelist
    from tenderagent.config import OfferTerms, Requisites
    from tenderagent.docs.fill import fill_docx, offer_lines
    from tenderagent.matching import PriceIndex, match_positions
    from tenderagent.models import Position

    p = tmp_path / "price.xlsx"
    make_price(p)
    items, _ = pricelist.load_items(p)
    req = Requisites(full_name="ООО «Ромашка»", unp="190000001", legal_address="г. Минск, ул. Ленина, 1",
                     bank_account="BY00ALFA30120000000000000000", bank_name="ЗАО «Альфа-Банк»")
    ms = match_positions([Position("Бумага офисная А4", 50), Position("Ручка шариковая синяя", 100),
                          Position("Маркер текстовыделитель", 10)], PriceIndex(items), OfferTerms(), 60)
    lines = offer_lines(ms, OfferTerms(), req)
    src = tmp_path / "form.docx"
    make_offer_form(src)
    dst = tmp_path / "out.docx"
    st = fill_docx(src, dst, lines, req, OfferTerms())
    assert st["offer_table"] and st["rows"] == 2
    d = docx.Document(dst)
    text = "\n".join(p.text for p in d.paragraphs)
    assert "ООО «Ромашка»" in text and "190000001" in text
    rows = [[c.text for c in r.cells] for r in d.tables[0].rows]
    assert rows[1][4] == "12,50" and rows[1][5] == "625,00" and rows[1][6] == "750,00"
    assert rows[2][4] == "0,85"
    assert rows[3][4] == ""  # маркера нет в прайсе
    assert rows[4][5] == "710,00"
    anketa = [[c.text for c in r.cells] for r in d.tables[1].rows]
    assert anketa[0][1] == "г. Минск, ул. Ленина, 1"
    assert "Альфа" in anketa[1][1]


LIST_HTML = """
<html><body><table id="auctions-list">
<tr><th>Название</th><th>Заказчик</th><th>Номер</th><th>Стоимость</th><th>Предложения до</th></tr>
<tr><td><a href="/tenders/all/view/1234567">Поставка бумаги офисной</a></td><td>ОАО «Завод»</td>
<td>auc0001234567</td><td>5 000,00 BYN</td><td>15.10.2099 10:00</td></tr>
<tr><td><a href="/tenders/all/view/1234568">Ремонт кровли</a></td><td>КУП</td>
<td>auc0001234568</td><td>90 000,00 BYN</td><td>16.10.2099</td></tr>
</table></body></html>
"""

DETAIL_HTML = """
<html><body><h1>Процедура закупки</h1>
<table class="details_table">
<tr><td>Номер процедуры</td><td>auc0001234567</td></tr>
<tr><td>Краткое описание предмета закупки</td><td>Поставка бумаги офисной</td></tr>
<tr><td>Наименование организации</td><td>ОАО «Завод», УНП 100000000</td></tr>
<tr><td>Дата и время окончания приема предложений</td><td>15.10.2099 10:00</td></tr>
<tr><td>Общая ориентировочная стоимость закупки</td><td>5 000,00 BYN</td></tr>
</table>
<table><tr><th>№ лота</th><th>Предмет закупки</th><th>Количество</th><th>Ед. изм.</th><th>Код ОКРБ</th></tr>
<tr><td>1</td><td>Бумага офисная А4</td><td>50</td><td>пачка</td><td>17.23.14.500</td></tr>
<tr><td>2</td><td>Ручка шариковая синяя</td><td>100</td><td>шт</td><td>32.99.12.100</td></tr></table>
<a href="/files/12345/Документация.docx">Документация.docx</a>
<a href="https://icetrade.by/info/files/777">Форма предложения.xlsx</a>
</body></html>
"""


def test_site_parsing():
    from tenderagent.models import Tender
    from tenderagent.sites import Http, load_site_configs
    from tenderagent.sites.base import GenericSite, tender_matches

    site = GenericSite(load_site_configs()["icetrade"], Http())
    ts = site.parse_list(LIST_HTML, "https://icetrade.by/search/auctions")
    assert [t.ext_id for t in ts] == ["1234567", "1234568"]
    assert ts[0].title == "Поставка бумаги офисной"
    assert ts[0].deadline == "2099-10-15T10:00"
    assert tender_matches(ts[0], ["бумага"], [], []) == "бумага"
    assert tender_matches(ts[1], ["бумага"], [], []) == ""
    t = Tender("icetrade", "1234567", "https://icetrade.by/tenders/all/view/1234567")
    site.parse_details(t, DETAIL_HTML, t.url)
    assert t.number == "auc0001234567"
    assert t.estimate == 5000.0
    assert t.customer.startswith("ОАО «Завод»")
    assert len(t.positions) == 2 and t.positions[0].qty == 50
    assert "17.23.14.500" in t.okrb
    assert len(t.documents) == 2
    assert tender_matches(t, [], ["17.23"], []) == "17.23"


def test_full_prepare(tmp_path):
    from tenderagent import pricelist
    from tenderagent.config import Requisites, Settings
    from tenderagent.db import DB, STATUS_READY
    from tenderagent.models import Position, Tender
    from tenderagent.pipeline import Engine, tender_folder

    p = tmp_path / "price.xlsx"
    make_price(p)
    s = Settings()
    s.requisites = Requisites(full_name="ООО «Ромашка»", unp="190000001")
    db = DB()
    eng = Engine(s, db, log=lambda *a: None)
    eng.set_price_items(pricelist.load_items(p)[0])
    t = Tender("manual", "x1", "https://example", title="Канцтовары", number="T-1", estimate=2000.0,
               positions=[Position("Канцтовары", 1)])
    folder = tender_folder(t)
    db.save_tender(t, folder=str(folder))
    (folder / "1. Документы заказчика").mkdir()
    make_offer_form(folder / "1. Документы заказчика" / "Приложение 2 Форма предложения.docx")
    a = eng.analyze(t, download=False)
    assert a.summary.total == 3 and a.summary.found == 2
    s.monitor.min_match_percent = 60
    assert eng.evaluate(t, a.matches).fits
    out = eng.prepare(t, a)
    names = {x.name for x in (out / "2. Наше предложение (для подачи)").iterdir()}
    assert "Приложение 2 Форма предложения (заполнено).docx" in names
    assert {"Сопроводительное письмо.docx", "Опись документов.docx", "Сведения об участнике.docx"} <= names
    assert (out / "3. Для проверки" / "ЧЕК-ЛИСТ перед подачей.docx").exists()
    assert (out / "Сведения о закупке.txt").exists()
    assert db.get(t.uid)["status"] == STATUS_READY


def test_butb_registry():
    """Реестр zakupki.butb.by (JSF): закупки без ссылок, разбираем таблицу."""
    from tenderagent.sites import Http, make_sites

    site = make_sites(Http())["butb"]
    html = (Path(__file__).parent / "fixtures" / "butb_reestr.html").read_text("utf-8")
    ts = site.parse_list(html, "https://zakupki.butb.by/auctions/reestrauctions.html;jsessionid=X")
    assert len(ts) >= 15
    t = ts[0]
    assert t.ext_id == "AU20261001394578"
    assert "Маска кислородная" in t.title
    assert t.estimate == 990.0
    assert t.deadline == "2026-10-13T23:59"
    assert "Медтехника" in t.customer
    assert t.fields["_butb_link"] == "fra:reestrAu:0:_t347"
    assert t.url == "https://zakupki.butb.by/auctions/reestrauctions.html"
    chesnok = [x for x in ts if "Чеснок" in x.title]
    assert chesnok and chesnok[0].estimate is None  # у запроса цен стоимость не указана
    assert chesnok[0].procedure == "заявка о ценах (тарифах)"


def test_goszakupki_real_pages():
    """Настоящие страницы goszakupki.by (сохранены программой в «Диагностика»)."""
    from tenderagent.models import Tender
    from tenderagent.sites import Http, make_sites

    site = make_sites(Http())["goszakupki"]
    fx = Path(__file__).parent / "fixtures"
    ts = site.parse_list((fx / "goszakupki_list.html").read_text("utf-8"), "https://goszakupki.by/tenders/posted")
    assert len(ts) == 20
    first = ts[0]
    assert first.ext_id == "single-source/view/3717571"
    assert first.estimate == 1620.0          # не «20261620» — год из даты не прилипает
    assert first.deadline == "2026-10-05T23:59"

    t = Tender("goszakupki", first.ext_id, "https://goszakupki.by/single-source/view/3717571")
    site.parse_details(t, (fx / "goszakupki_card.html").read_text("utf-8"), t.url)
    assert t.number == "auc0003717571"
    assert t.customer == "Докшицкий районный исполнительный комитет" and t.customer_unp == "300013863"
    assert t.deadline == "2026-10-05T23:59"
    p = t.positions[0]
    assert (p.name, p.qty, p.unit, p.price_limit) == ("Бумага офисная А4", 150.0, "пачка", 1620.0)
    names = [d.name for d in t.documents]
    assert len(names) == 3 and not any("Регламент" in n for n in names)


def test_requisites_from_company_card(tmp_path):
    """Карточка предприятия в свободной форме → поля реквизитов."""
    import docx

    from tenderagent.requisites_import import parse_requisites, read_text

    d = docx.Document()
    for line in [
        "Общество с ограниченной ответственностью «Ромашка»",
        "УНП 190000001",
        "220000, Минская область, г.Минск, ул.Ленина, 1 «а» пом. 2",
        "р/сч  №BY00ALFA30120000000000000000 ЗАО «Альфа-Банк», г.Минск, БИК ALFABY2X",
        "GLN основной 4810000000001 220000, Беларусь, г.Минск, ул. Ленина, д. 1 А",
        "Директор Иванов Иван Иванович",
        "На основании Устава",
        "Электронная почта info@romashka.by",
        "Тел. +375 (17) 123-45-67",
    ]:
        d.add_paragraph(line)
    p = tmp_path / "card.docx"
    d.save(p)
    r = parse_requisites(read_text(p))
    assert r["full_name"] == "Общество с ограниченной ответственностью «Ромашка»"
    assert r["short_name"] == "ООО «Ромашка»"
    assert r["unp"] == "190000001"
    assert r["legal_address"].startswith("220000, Минская область")
    assert r["bank_account"] == "BY00ALFA30120000000000000000"
    assert r["bank_name"] == "ЗАО «Альфа-Банк»" and r["bank_address"] == "г.Минск"
    assert r["bank_bic"] == "ALFABY2X"
    assert (r["director_position"], r["director_name"], r["director_short"]) == (
        "Директор", "Иванов Иван Иванович", "И.И. Иванов")
    assert r["acts_on"] == "Устава"
    assert r["email"] == "info@romashka.by"
    assert r["phone"] == "+375 (17) 123-45-67"  # не кусок номера счёта и не GLN


def test_tyre_pricelist_like_belshina(tmp_path):
    """Прайскурант вида «Белшины»: «Номенклатурный номер» — это код, а не наименование;
    «А/ШИНА 215/75R16C» должна находиться по «Автошина 215/75 R16C»."""
    import openpyxl

    from tenderagent import pricelist, textnorm
    from tenderagent.matching import PriceIndex

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["ПРЕЙСКУРАНТ № 5"])
    ws.append(["Номенклатурный\nномер", "Код\nбазы ГП", "Наименование продукции", "Дата\nустановления\nцены",
               "Цена,\nBYN", "Штрихкод"])
    ws.append(["ШИНЫ"])
    rows = [("0101", "А/ШИНА 215/75R16C BEL-313 СЕР Л/ГРУЗ", 218.5),
            ("0102", "А/ШИНА 205/55R16 BEL-262 СЕР ЛЕГК Б/К", 163.02),
            ("0103", "А/ШИНА 225/65R17 BEL-295 СЕР ЛЕГК Б/К", 224.84),
            ("0104", "А/ШИНА 185/75R16C BEL-293 СЕР Л/ГР", 180.49)]
    for code, name, price in rows:
        ws.append([code, "K1", name, "01.05.2026", price, 4811644000000])
    p = tmp_path / "price.xlsx"
    wb.save(p)
    items, table = pricelist.load_items(p)
    assert pricelist.guess_columns(table.headers)["name"] == "Наименование продукции"
    assert items[0].name.startswith("А/ШИНА") and items[0].code == "0101"
    idx = PriceIndex(items)
    best = lambda q: idx.candidates(q, 1)[0]  # noqa: E731
    assert best("Автошина 215/75 R16C")[0].code == "0101" and best("Автошина 215/75 R16C")[1] >= 80
    assert best("Шина 205/55 R16 зимняя")[0].code == "0102"
    assert best("Автошины 225/65 R17 (зима)")[0].code == "0103"
    assert best("Закупка зимних автошин 285/65 R16C")[1] < 70   # такого размера нет
    # лишние слова не мешают, если типоразмер совпал
    gaz = best("Автошины на автомобиль ГАЗ 330273, 185/75R16C")
    assert gaz[0].code == "0104" and gaz[1] >= 85
    assert textnorm.contains_keyword("Автошины зимние", "шины")


def test_download_rejects_html_and_fixes_extension(tmp_path):
    from tenderagent.sites.base import _fix_extension, _looks_like_html

    assert _looks_like_html(b"<!DOCTYPE html><html><body>login</body></html>")
    assert not _looks_like_html(b"%PDF-1.7 ...")
    assert _fix_extension("прейскуранту", b"\xd0\xcf\x11\xe0" + b"\x00" * 100) == "прейскуранту.doc"
    assert _fix_extension("tz.pdf", b"%PDF-1.4") == "tz.pdf"
    assert _fix_extension("dogovor", b"PK\x03\x04....word/document.xml") == "dogovor.docx"


def test_folders_by_customer(tmp_path):
    """Тендеры / <Заказчик> / <дата № номер — предмет>; старые папки переносятся."""
    from tenderagent.config import tenders_dir
    from tenderagent.folders import DOCS, OFFER, customer_dir_name, tender_folder
    from tenderagent.models import Document, Tender

    assert customer_dir_name('Республиканское дочернее торговое унитарное предприятие "Медтехника" г.Гомель') \
        == "РДТУП Медтехника г.Гомель"
    assert customer_dir_name("Докшицкий районный исполнительный комитет") == "Докшицкий райисполком"
    assert customer_dir_name("") == "Заказчик не указан"

    t = Tender("goszakupki", "single-source/view/1", "u", title="Бумага офисная А4", number="auc0003717571",
               customer="Докшицкий районный исполнительный комитет", published="2026-10-02T00:00")
    f = tender_folder(t)
    assert f.parent == tenders_dir() / "Докшицкий райисполком"
    assert f.name == "2026-10-02 № auc0003717571 Бумага офисная А4"
    # второй тендер того же заказчика — в той же папке заказчика
    t2 = Tender("icetrade", "2", "u", title="Ручки", number="auc2", customer=t.customer)
    assert tender_folder(t2).parent == f.parent

    # папка старого формата (1.0.x): Тендеры/<дата площадка номер предмет>/Документация
    old = tenders_dir() / "2026-10-06 goszakupki 123 Шины"
    (old / "Документация").mkdir(parents=True)
    (old / "Документация" / "tz.pdf").write_bytes(b"%PDF")
    (old / "Для подачи").mkdir()
    t3 = Tender("goszakupki", "3", "u", title="Шины", number="123", customer="ОАО «Завод»",
                documents=[Document("tz.pdf", "u", str(old / "Документация" / "tz.pdf"))])
    new = tender_folder(t3, str(old))
    assert not old.exists() and new.parent.name == "ОАО Завод"
    assert (new / DOCS / "tz.pdf").exists() and (new / OFFER).exists()
    assert t3.documents[0].local_path == str(new / DOCS / "tz.pdf")


def test_number_not_taken_from_contact_name():
    from tenderagent.models import Tender
    from tenderagent.sites import Http, make_sites

    html = """<html><body><h1>Процедура закупки</h1><table>
    <tr><td>Номер телефона / контактное лицо №</td><td>Ахрамович Василина Анатольевна</td></tr>
    <tr><td>Предмет закупки</td><td>Закупка автошин</td></tr></table></body></html>"""
    site = make_sites(Http())["goszakupki"]
    t = Tender("goszakupki", "tender/view/3720000", "https://goszakupki.by/tender/view/3720000")
    site.parse_details(t, html, t.url)
    assert "Ахрамович" not in t.number


def test_signed_documents_are_unwrapped(tmp_path):
    """goszakupki отдаёт документы в «конверте» ЭЦП (CMS) — Word пишет «содержимое не удалось прочитать».
    Программа достаёт исходный документ, подписанный оригинал кладёт отдельно."""
    import shutil

    from tenderagent.docs.extract import analyze_folder
    from tenderagent.docs.unwrap import SIGNED_DIR, unwrap

    signed = (Path(__file__).parent / "fixtures" / "signed_cms.docx").read_bytes()
    assert signed[:2] == b"\x30\x82"
    inner = unwrap(signed)
    assert inner[:4] == b"PK\x03\x04"
    # PDF в конверте: начало по сигнатуре, хвост подписи отрезается
    fake = b"\x30\x82\x10\x00" + b"\x06\x09" + b"x" * 40 + b"%PDF-1.4 body %%EOF" + b"\xa0\x82signature"
    assert unwrap(fake) == b"%PDF-1.4 body %%EOF"
    assert unwrap(b"%PDF-1.4 plain") == b"%PDF-1.4 plain"

    folder = tmp_path / "docs"
    folder.mkdir()
    shutil.copy(Path(__file__).parent / "fixtures" / "signed_cms.docx", folder / "zapros-ceny-avtoshiny-zimnie")
    infos = analyze_folder(folder, log=lambda *a: None)
    assert [i.path.name for i in infos] == ["zapros-ceny-avtoshiny-zimnie.docx"]
    assert infos[0].positions[0].name == "Автошина 215/75 R16C"
    assert (folder / SIGNED_DIR / "zapros-ceny-avtoshiny-zimnie.p7s").exists()


def test_paths_fit_windows_limit():
    """Excel/Word не открывают файлы с путём длиннее 259 символов."""
    from tenderagent.folders import CHECK, OFFER, customer_dir_name, tender_dir_name
    from tenderagent.models import Tender
    from tenderagent.sites.base import _safe_name

    t = Tender("goszakupki", "x", "u", number="auc0003799999",
               title="Закупка зимних автошин 285/65 R16C 131 R M+S для автомобилей филиала " * 3,
               customer="Государственное учреждение образования «Средняя школа № 1 имени героя "
                        "Советского Союза И.И. Иванова д.Вельямовичи Брестского района», УНП 200000000",
               published="2026-10-06T00:00")
    base = r"C:\Users\user\Documents\ТендерАгент\Тендеры"
    longest_file = _safe_name("Приложение № 2 к документации о закупке — форма ценового предложения участника "
                              "(заполнено).xlsx")
    for sub in (OFFER, CHECK):
        path = "\\".join([base, customer_dir_name(t.customer), tender_dir_name(t), sub, longest_file])
        assert len(path) <= 250, (len(path), path)
    assert longest_file.endswith(".xlsx")


def test_goszakupki_file_link_returns_info_then_file(tmp_path):
    """goszakupki: ссылка на документ отдаёт JSON-справку, сам файл — по той же ссылке с &download=1."""
    from tenderagent.sites.base import Http

    docx_bytes = (Path(__file__).parent / "fixtures" / "signed_cms.docx").read_bytes()
    stub = ('{"ok":true,"info":{"name":"proekt-dogovora-prilozhenie-1_1791358415.docx",'
            '"size":"24.67Кб","created":"07.10.2026","key":1}}').encode()

    class Resp:
        def __init__(self, content, url):
            self.content, self.url, self.headers, self.ok = content, url, {}, True

        def raise_for_status(self):
            pass

    calls = []

    class Sess:
        headers = {}

        def get(self, url, **kw):
            calls.append(url)
            return Resp(docx_bytes if url.endswith("download=1") else stub, url)

    h = Http(log=lambda *a: None)
    h.s = Sess()
    url = "https://goszakupki.by/single-source/get-file/3717571?c=detail&f=1"
    p = h.download(url, tmp_path, "proekt-dogovora")
    assert calls == [url, url + "&download=1"]
    assert p.name == "proekt-dogovora-prilozhenie-1_1791358415.docx"
    assert p.read_bytes()[:4] == b"PK\x03\x04"   # настоящий документ (и без конверта ЭЦП)


def test_unit_prices_and_markup_fit():
    """Цена заказчика за единицу, расхождение и подбор наценки под ориентир (пример из goszakupki:
    24 шины 225/75 R16C, ориентир 6 035,90 BYN с НДС)."""
    from tenderagent.config import OfferTerms
    from tenderagent.matching import customer_unit_price, markup_for_target, unit_price
    from tenderagent.models import Match, Position, PriceItem

    item = PriceItem("А/ШИНА 225/75R16C BEL-500 СЕР Л/ГР Б/К", 244.19)
    m = Match(Position("Автошина зимняя 225/75 R 16C 121/120 R", qty=24, unit="штук"), item, 85)
    assert customer_unit_price(m, 1, 6035.90) == 251.5          # 6035,90 / 24
    # прайс с НДС: наша цена с НДС 244,19 → на 2,9 % ниже заказчика
    terms = OfferTerms(prices_include_vat=True)
    m.price = unit_price(item, terms)
    assert round(m.price * 1.2, 2) == 244.19
    # подобрать наценку, чтобы быть на 1 % ниже ориентира
    mk = markup_for_target([m], terms, 1.2, 6035.90 * 0.99, {})
    terms.markup_percent = int(mk * 10) / 10
    total = unit_price(item, terms) * 24 * 1.2
    assert total <= 6035.90 * 0.99 and total > 6035.90 * 0.98


def test_cleanup_expired(tmp_path):
    """«Очистить неактуальное»: просроченные удаляются (с папками по желанию), поданные остаются,
    удалённые не возвращаются при следующем мониторинге."""
    from tenderagent.config import Settings
    from tenderagent.db import DB
    from tenderagent.models import Tender
    from tenderagent.pipeline import Engine, tender_folder

    db = DB()
    eng = Engine(Settings(), db, log=lambda *a: None)
    old = Tender("goszakupki", "1", "u", title="Шины", number="a1", customer="ОАО Завод", deadline="2020-01-01T10:00")
    sub = Tender("goszakupki", "2", "u", title="Бумага", number="a2", customer="ОАО Завод", deadline="2020-01-01T10:00")
    new = Tender("goszakupki", "3", "u", title="Ручки", number="a3", customer="КУП", deadline="2099-01-01T10:00")
    for t in (old, sub, new):
        db.save_tender(t, folder=str(tender_folder(t)))
    db.set_status(sub.uid, "Подан")
    assert eng.mark_expired() == 1
    victims = eng.cleanup_candidates()
    assert [r["uid"] for r in victims] == [old.uid]
    folder = Path(victims[0]["folder"])
    assert eng.cleanup(victims, delete_folders=True) == (1, 1)
    assert not folder.exists() and folder.parent.exists()     # у заказчика остался поданный тендер
    assert db.get(old.uid) is None and db.has(old.uid)        # мониторинг его не добавит снова
    assert db.get(sub.uid)["status"] == "Подан" and db.get(new.uid) is not None


def test_customer_terms_docs_and_marking(tmp_path):
    """Заказчик — не контактное лицо; ссылки «get-file» без расширения — документы;
    сроки поставки/оплаты берутся у заказчика; справка о маркировке — если её требуют."""
    from tenderagent.config import OfferTerms
    from tenderagent.models import Tender
    from tenderagent.pipeline import customer_terms
    from tenderagent.sites import Http, make_sites

    html = """<html><body><h1>Процедура закупки № auc0003722747</h1><table>
    <tr><td>Контактное лицо заказчика</td><td>Сушко Анастасия Васильевна</td></tr>
    <tr><td>Наименование заказчика</td><td>КУП «Горводоканал»</td></tr>
    <tr><td>Предмет закупки</td><td>Автошины 225/75R16C</td></tr>
    <tr><td>Срок поставки</td><td>в течение 5 рабочих дней с даты заявки</td></tr>
    <tr><td>Условия оплаты</td><td>Отсрочка платежа 30 календарных дней</td></tr>
    <tr><td>Требования к маркировке</td><td>товар должен быть маркирован средствами идентификации</td></tr>
    </table>
    <a class="modal-link" href="/marketing/get-file/3722747?c=detail&f=0">Проект договора</a>
    <a class="modal-link" href="/marketing/get-file/3722747?c=detail&f=1">Предложение о закупке</a>
    </body></html>"""
    site = make_sites(Http())["goszakupki"]
    t = Tender("goszakupki", "marketing/view/3722747", "https://goszakupki.by/marketing/view/3722747")
    site.parse_details(t, html, t.url)
    assert t.customer == "КУП «Горводоканал»" and t.number == "auc0003722747"
    assert [d.name for d in t.documents] == ["Проект договора", "Предложение о закупке"]
    terms, taken = customer_terms(OfferTerms(), t.fields, {})
    assert terms.delivery_term == "в течение 5 рабочих дней с даты заявки"
    assert terms.payment_terms == "отсрочка платежа 30 календарных дней"
    assert "Срок поставки" in taken and t.fields["Требования к маркировке"]


def test_offer_uses_our_product_name_and_valid_number(tmp_path):
    """В нашем ценовом предложении — наименование из нашего прайса; номер процедуры — не ФИО."""
    import docx

    from tenderagent.config import OfferTerms, Requisites
    from tenderagent.docs.fill import offer_lines
    from tenderagent.docs.generate import offer_docx
    from tenderagent.models import Match, Position, PriceItem, Tender

    t = Tender("goszakupki", "marketing/view/3725000", "u", title="Шины 8,25 R20",
               number="Мамай Илона Николаевна, +375177124714")
    assert t.num == "3725000"
    assert Tender("x", "1", "u", number="auc0003722747").num == "auc0003722747"
    m = Match(Position("Шины 8,25 R20 (Вилейский центр)", 12, "штук"),
              PriceItem("А/ШИНА 8.25R20 К-84МБ,У-2", 555.67), 90, price=555.67)
    p = tmp_path / "offer.docx"
    offer_docx(p, t, offer_lines([m], OfferTerms(), Requisites(full_name="ООО «Миртаер»")),
               Requisites(full_name="ООО «Миртаер»"), OfferTerms())
    d = docx.Document(p)
    assert d.tables[0].rows[1].cells[1].text == "А/ШИНА 8.25R20 К-84МБ,У-2"
    text = "\n".join(x.text for x in d.paragraphs)
    assert "№ 3725000" in text and "Мамай" not in text


def test_old_stub_files_are_replaced(tmp_path):
    """Старая справка «{"ok":true,"info":…}» под именем .doc удаляется, настоящий файл
    ложится на её место (а не рядом с «(1)» в имени)."""
    from tenderagent.docs.extract import purge_bad_downloads
    from tenderagent.sites.base import Http

    stub = b'{"ok":true,"info":{"name":"zayavka-na-predelnuju-5_1791295129.doc","size":"66.00\xd0\x9a\xd0\xb1","key":1}}'
    real = b"\xd0\xcf\x11\xe0" + b"\x00" * 600
    (tmp_path / "zayavka-na-predelnuju-5_1791295129.doc").write_bytes(stub)
    (tmp_path / "ok.pdf").write_bytes(b"%PDF-1.4 real")

    class Resp:
        def __init__(self, content):
            self.content, self.url, self.headers, self.ok = content, "u", {}, True

        def raise_for_status(self):
            pass

    class Sess:
        headers = {}

        def get(self, url, **kw):
            return Resp(real if url.endswith("download=1") else stub)

    h = Http(log=lambda *a: None)
    h.s = Sess()
    p = h.download("https://goszakupki.by/marketing/get-file/1?c=detail&f=1", tmp_path, "x")
    assert p.name == "zayavka-na-predelnuju-5_1791295129.doc" and p.read_bytes() == real
    assert not (tmp_path / "zayavka-na-predelnuju-5_1791295129 (1).doc").exists()

    (tmp_path / "old.doc").write_bytes(stub)
    assert purge_bad_downloads(tmp_path, log=lambda *a: None) == 1
    assert not (tmp_path / "old.doc").exists() and (tmp_path / "ok.pdf").exists()


def test_goszakupki_zip_fallback_and_marketing_page(tmp_path):
    """Если &download=1 не отдал файл, берём «Получить архив» (&downloadZip=1):
    в архиве документ и подпись .sgn. Плюс разбор страницы «заявки о ценах»."""
    import io
    import zipfile

    from tenderagent.docs.unwrap import SIGNED_DIR
    from tenderagent.models import Tender
    from tenderagent.sites import Http, make_sites

    stub = b'{"ok":true,"info":{"name":"zayavka-na-predelnuju-5_1791295129.doc","size":"66.00Kb","key":1}}'
    doc = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 800
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("zayavka-na-predelnuju-5_1791295129.doc", doc)
        z.writestr("zayavka-na-predelnuju-5_1791295129.doc.sgn", b"\x30\x82\x0f\x87signature")

    class Resp:
        def __init__(self, content):
            self.content, self.url, self.headers, self.ok = content, "u", {}, True

        def raise_for_status(self):
            pass

    calls = []

    class Sess:
        headers = {}

        def get(self, url, **kw):
            calls.append(url.rsplit("&", 1)[-1])
            return Resp(buf.getvalue() if url.endswith("downloadZip=1") else stub)

    h = Http(log=lambda *a: None)
    h.s = Sess()
    p = h.download("https://goszakupki.by/marketing/get-file/3728554?c=detail&f=1", tmp_path, "x")
    assert calls[-2:] == ["download=1", "downloadZip=1"]
    assert p.name == "zayavka-na-predelnuju-5_1791295129.doc" and p.read_bytes() == doc
    assert (tmp_path / SIGNED_DIR / "zayavka-na-predelnuju-5_1791295129.doc.sgn").exists()

    html = (Path(__file__).parent / "fixtures" / "goszakupki_marketing.html").read_text("utf-8")
    t = Tender("goszakupki", "marketing/view/3728554", "https://goszakupki.by/marketing/view/3728554")
    make_sites(Http())["goszakupki"].parse_details(t, html, t.url)
    assert t.number == "auc0003728554"
    assert t.customer == 'Учреждение здравоохранения "Жлобинская центральная районная больница"'
    assert [(p.name, p.qty) for p in t.positions][2] == ("Автошина зимняя 235/65R16С", 4.0)
    assert [d.name for d in t.documents] == ["zad.-na-zak_1791295123.pdf", "zayavka-na-predelnuju-5_1791295129.doc"]
    assert t.fields["Условия оплаты"] == "Согласно договора"
    assert t.fields["Срок поставки"] == "c 20.10.2026 по 26.10.2026"
    assert t.fields["Место поставки"].startswith("Республика Беларусь, Гомельская область, 247210")


def test_updater_check_and_script(tmp_path, monkeypatch):
    """Обновление: находит новый выпуск ta-v<версия>, скачивает .exe и пишет скрипт подмены."""
    import sys as _sys

    from tenderagent import updater

    assert updater.is_newer("1.1.1", "1.1.0") and not updater.is_newer("1.0.18", "1.1.0")

    class R:
        def __init__(self, js=None, content=b""):
            self._js, self.content, self.headers = js, content, {"content-length": str(len(content))}

        def raise_for_status(self):
            pass

        def json(self):
            return self._js

        def iter_content(self, n):
            yield self.content

        def __enter__(self):
            return self

        def __exit__(self, *a):
            return False

    release = {"tag_name": "ta-v9.0.0", "body": "Что нового", "html_url": "https://github.com/x",
               "assets": [{"name": "TenderAgent.exe", "browser_download_url": "https://dl/exe"},
                          {"name": "TenderAgent-folder.zip", "browser_download_url": "https://dl/zip"}]}
    monkeypatch.setattr(updater.requests, "get",
                        lambda url, **kw: R(release) if "api.github.com" in url else R(content=b"MZnewexe"))
    rel = updater.check()
    assert rel.version == "9.0.0" and rel.exe_url == "https://dl/exe" and rel.notes == "Что нового"
    release["tag_name"] = "tender-agent-v1.0.99"   # старый формат тегов — не предлагать
    assert updater.check() is None
    release["tag_name"] = "ta-v9.0.0"

    exe = tmp_path / "ТендерАгент.exe"
    exe.write_bytes(b"MZold")
    monkeypatch.setattr(_sys, "frozen", True, raising=False)
    monkeypatch.setattr(_sys, "executable", str(exe))
    monkeypatch.setattr(_sys, "_MEIPASS", str(tmp_path / "_MEI123"), raising=False)
    monkeypatch.setattr(updater.tempfile, "mkdtemp", lambda prefix="": str(tmp_path / "upd"))
    (tmp_path / "upd").mkdir()
    started = []
    monkeypatch.setattr(updater.subprocess, "Popen", lambda args, **kw: started.append(args))
    assert updater.install_kind() == "onefile"
    updater.apply_and_restart(rel)
    script = (tmp_path / "upd" / "update.ps1").read_text("utf-8-sig")
    assert (tmp_path / "upd" / "new.exe").read_bytes() == b"MZnewexe"
    assert "ТендерАгент.exe" in script and "Wait-Process" in script and "Start-Process" in script
    assert started and started[0][0] == "powershell"
